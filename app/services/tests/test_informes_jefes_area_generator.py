import os
import tempfile
import unittest

import openpyxl
import pandas as pd

from app.services.etl_jobs.informes_jefes_area_generator import (
    COLUMNS_WIDE,
    HOJA_WIDE,
    _construir_base_desde_fuentes,
    generate_informes_jefes_area,
)

SUB_GRUPOS = [
    ("PRODUCCION", "PRODUCCION TOTAL", "UND"),
    ("PRODUCCION", "PRODUCCION EXPORTABLE", "UND"),
    ("PRODUCCION", "EXPORTADO", "UND"),
    ("AREAS", "HECTAREAS", "UND"),
    ("PLANTAS", "PLANTAS SEMBRADAS", "UND"),
    ("PLANTAS", "PLANTAS EN PRODUCCION", "UND"),
    ("PERSONAL", "PERSONAS + HORAS EXTRAS", "UND"),
    ("NOMINA", "NOMINA + EXTRAS + PREST/SOCIALES/ISS", "$"),
    ("COMPRAS", "COMPRAS $ GASTOS", "$"),
    ("TRANSPORTES", "COSTO TRANSPORTE PERSONAL", "$"),
    ("TRANSPORTES", "TRANSPORTE FLOR AEROPUERTO", "$"),
]

FINCAS = [("PALMAS", "CULTIVO", 100), ("PALERMO", "CULTIVO", 200)]
TIPOS_AÑOS = [("REAL", 2025), ("REAL", 2026), ("TEORICO", 2026)]
AREAS_ORDEN = ["CULTIVO", "Empresa", "PC"]
MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]


def _build_base_datos(path: str) -> None:
    wb = openpyxl.Workbook()

    ws_base = wb.active
    ws_base.title = "base"
    ws_base.append([
        "FINCA", "AREA", "GFLOR", "FLOR", "GRUPO", "SUB_GRUPO", "TIPO", "UMEDIDA", "AÑO", "MES", "Valor",
    ])
    for finca, area, valor in FINCAS:
        for grupo, sub_grupo, umedida in SUB_GRUPOS:
            for tipo, año in TIPOS_AÑOS:
                for mes in MESES:
                    ws_base.append([
                        finca, area, "CLAVEL", "CLAVEL", grupo, sub_grupo, tipo, umedida, año, mes, valor,
                    ])

    ws_orden = wb.create_sheet("orden")
    ws_orden.append(["sub_grupo", "area", "orden", "opera"])
    orden_n = 1
    for area in AREAS_ORDEN:
        for _grupo, sub_grupo, _umedida in SUB_GRUPOS:
            ws_orden.append([sub_grupo, area, orden_n, "suma"])
            orden_n += 1

    wb.save(path)


class InformesJefesAreaGeneratorTests(unittest.TestCase):
    def test_generates_workbook_with_expected_sheets_and_totals(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = os.path.join(tmp_dir, "base_datos.xlsx")
            out_dir = os.path.join(tmp_dir, "out")
            _build_base_datos(input_path)

            result = generate_informes_jefes_area({
                "file": input_path, "out_dir": out_dir, "use_live_source": False,
            })

            self.assertTrue(result.startswith("xlsx:"))
            output_path = result.split("xlsx:", 1)[1]
            self.assertTrue(os.path.isfile(output_path))

            wb = openpyxl.load_workbook(output_path, data_only=True)
            self.assertEqual(
                wb.sheetnames,
                [
                    "Informes Consolidado",
                    "Informes Palmas",
                    "Informes Palermo",
                    "Informes Propagacion",
                    "Informes PC",
                    "Consolidado",
                ],
            )

            # fila 1 en blanco, fila 2 con el rotulo, fila 3 encabezado, datos desde la fila 4
            ws = wb["Informes Palmas"]
            self.assertEqual(ws.cell(row=2, column=2).value, "JEFE PALMAS")
            self.assertEqual(ws.cell(row=3, column=1).value, "finca")

            def find_row(ws_, finca, area, grupo, sub_grupo, tipo, año):
                for r in range(4, ws_.max_row + 1):
                    if (
                        ws_.cell(row=r, column=1).value == finca
                        and ws_.cell(row=r, column=2).value == area
                        and ws_.cell(row=r, column=5).value == grupo
                        and ws_.cell(row=r, column=6).value == sub_grupo
                        and ws_.cell(row=r, column=7).value == tipo
                        and ws_.cell(row=r, column=8).value == año
                    ):
                        return r
                return None

            r = find_row(ws, "PALMAS", "CULTIVO", "PRODUCCION", "PRODUCCION TOTAL", "REAL", 2025)
            self.assertIsNotNone(r)
            self.assertEqual(ws.cell(row=r, column=10).value, 100)

            # 'Informes Consolidado' debe sumar Palmas (100) + Palermo (200)
            ws_total = wb["Informes Consolidado"]
            r_total = find_row(ws_total, "Todo", "Todo", "PRODUCCION", "PRODUCCION TOTAL", "REAL", 2025)
            self.assertIsNotNone(r_total)
            self.assertEqual(ws_total.cell(row=r_total, column=10).value, 300)


def _build_tabla_wide(path: str, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = HOJA_WIDE
    ws.append(COLUMNS_WIDE)
    for row in rows:
        ws.append([row.get(col, 0) for col in COLUMNS_WIDE])
    wb.save(path)


class ConstruirBaseDesdeFuentesTests(unittest.TestCase):
    """Cubre la reconstruccion en Python del flujo de Power Query
    (tablas<año> + datos_YYYY congelados -> 'base'), en particular el reparto
    de flores combinadas ('CLA Y ROS', etc.) usando la hoja 'distribuciones'."""

    def test_reparte_flor_combinada_segun_distribuciones_y_filtra_indices_y_anios(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tablas_dir = os.path.join(tmp_dir, "tablas2026")
            os.makedirs(tablas_dir)

            filas = [
                {
                    "id": 1, "AREA": "CULTIVO", "FINCA": "PALMAS", "GFLOR": "CLAVEL", "FLOR": "CLAVEL",
                    "AÑO": 2025, "SUB_GRUPO": "PRODUCCION TOTAL", "TIPO": "REAL", "GRUPO": "PRODUCCION",
                    "UMEDIDA": "UND", "ENERO": 500,
                },
                {
                    # fila combinada: se debe repartir entre CLAVEL y ROSAS COLORES
                    "id": 2, "AREA": "CULTIVO", "FINCA": "PALMAS", "GFLOR": "TODAS", "FLOR": "CLA Y ROS",
                    "AÑO": 2025, "SUB_GRUPO": "COMPRAS $ GASTOS", "TIPO": "REAL", "GRUPO": "COMPRAS",
                    "UMEDIDA": "$", "ENERO": 100,
                },
                {
                    # debe quedar excluida: GRUPO == 'INDICES'
                    "id": 3, "AREA": "CULTIVO", "FINCA": "PALMAS", "GFLOR": "CLAVEL", "FLOR": "CLAVEL",
                    "AÑO": 2025, "SUB_GRUPO": "ALGUN INDICE", "TIPO": "REAL", "GRUPO": "INDICES",
                    "UMEDIDA": "UND", "ENERO": 999,
                },
                {
                    # debe quedar excluida: AÑO en (2019, 2020)
                    "id": 4, "AREA": "CULTIVO", "FINCA": "PALMAS", "GFLOR": "CLAVEL", "FLOR": "CLAVEL",
                    "AÑO": 2019, "SUB_GRUPO": "PRODUCCION TOTAL", "TIPO": "REAL", "GRUPO": "PRODUCCION",
                    "UMEDIDA": "UND", "ENERO": 999,
                },
            ]
            _build_tabla_wide(os.path.join(tablas_dir, "produccion.xlsx"), filas)

            anios_congelados_paths = []
            for year in (2023, 2024, 2025):
                p = os.path.join(tmp_dir, f"datos_{year}.xlsx")
                _build_tabla_wide(p, [])
                anios_congelados_paths.append(p)

            dist_df = pd.DataFrame([{
                "FINCA": "PALMAS", "AREA": "CULTIVO", "GFLOR": "TODAS", "FLOR": "CLA Y ROS",
                "CLAVEL": 0.6, "MINIATURA": 0.0, "ROSAS COLORES": 0.4, "ROSAS ROJAS": 0.0,
            }])

            resultado = _construir_base_desde_fuentes(tablas_dir, anios_congelados_paths, dist_df)

            self.assertNotIn("INDICES", resultado["GRUPO"].unique())
            self.assertNotIn(2019, resultado["AÑO"].unique())
            self.assertNotIn("CLA Y ROS", resultado["FLOR"].unique())

            compras_enero = resultado[
                (resultado["FINCA"] == "PALMAS")
                & (resultado["SUB_GRUPO"] == "COMPRAS $ GASTOS")
                & (resultado["AÑO"] == 2025)
                & (resultado["MES"] == "ENERO")
            ]
            clavel = compras_enero[compras_enero["FLOR"] == "CLAVEL"]["Valor"].iloc[0]
            rosas_colores = compras_enero[compras_enero["FLOR"] == "ROSAS COLORES"]["Valor"].iloc[0]
            self.assertAlmostEqual(clavel, 60.0)
            self.assertAlmostEqual(rosas_colores, 40.0)

            produccion_enero = resultado[
                (resultado["SUB_GRUPO"] == "PRODUCCION TOTAL")
                & (resultado["AÑO"] == 2025)
                & (resultado["MES"] == "ENERO")
            ]
            self.assertEqual(len(produccion_enero), 1)
            self.assertEqual(produccion_enero["Valor"].iloc[0], 500)


if __name__ == "__main__":
    unittest.main()
