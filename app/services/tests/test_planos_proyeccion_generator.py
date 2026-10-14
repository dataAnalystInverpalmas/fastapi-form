import unittest
from pathlib import Path
from unittest.mock import patch

import mysql.connector
import pandas as pd

from app.services.etl_jobs.planos_proyeccion_generator import (
    HOJA1_COLUMNS,
    SOURCE_COLUMNS,
    _attach_id_lote_siembra,
    _build_ajustes_export,
    _build_hoja1_export,
    _build_hoja2_export,
    _build_subir_proyeccion_body,
    _build_tika_upload_body,
    _fill_subir_proyeccion_workbook,
    _hoja1_week_range,
    _load_base_from_staging,
    _load_quipus_excel,
    _notify_tika_ajuste,
    _transform_base_sheet,
    _transform_lotes_siembra,
    _write_ajustes_por_finca,
    generate_planos_proyeccion_hoja1,
)
from app.services.utils.date_service import get_next_week_thursday, yyww_series_to_date, yyww_to_date

# Corte anterior a todas las semanas de _build_datos_df, para no filtrar nada.
DESDE_TODO = pd.Timestamp("2025-01-01")


class PlanosProyeccionGeneratorTests(unittest.TestCase):
    def _build_raw(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "finca": [2.0, 2.0, None],
                "bloque": [1.0, 3.0, 1.0],
                "variedad": ["2048", "2149", "2048"],
                "cosecha": ["SE2638", "JL2631", "SE2638"],
                "nvari": ["FARIDA", "MOON LIGHT", "FARIDA"],
                "fec_siem": pd.to_datetime(["2026-03-11", "2026-01-21", "2026-03-11"]),
                "rev": [0.0, 0.0, 0.0],
                "amort": [1.0, 1.0, 1.0],
                "proyecta_quipus": [0.88, 1.17, 0.88],
                "edad": [26, 33, 26],
            }
        )

    def test_transform_extracts_and_normalizes_expected_columns(self) -> None:
        df, audit = _transform_base_sheet(self._build_raw())

        self.assertEqual(list(df.columns), [
            "finca", "bloque", "variedad", "cosecha", "nvari", "yyww_siembra", "rev", "amort", "proyecta_quipus",
        ])
        self.assertEqual(len(df), 2)
        self.assertEqual(audit["raw_rows"], 3)
        self.assertEqual(audit["rows_after_dropna_finca"], 2)
        self.assertEqual(audit["final_rows"], 2)

        row = df.iloc[0]
        self.assertEqual(row["finca"], "002")
        self.assertEqual(row["bloque"], "1")
        self.assertEqual(row["variedad"], "2048")
        self.assertEqual(row["nvari"], "FARIDA")
        self.assertEqual(row["yyww_siembra"], "2611")

    def test_transform_replaces_spaces_in_nvari_with_underscore(self) -> None:
        df, _audit = _transform_base_sheet(self._build_raw())
        self.assertEqual(df.iloc[1]["nvari"], "MOON_LIGHT")

    def test_transform_raises_when_required_column_missing(self) -> None:
        df_raw = self._build_raw().drop(columns=["proyecta_quipus"])
        with self.assertRaises(ValueError):
            _transform_base_sheet(df_raw)

    def test_transform_keeps_variedad_as_text_without_numeric_coercion(self) -> None:
        df_raw = self._build_raw()
        df_raw["variedad"] = ["A048", "2149", "2048"]

        df, _audit = _transform_base_sheet(df_raw)

        self.assertEqual(df.iloc[0]["variedad"], "A048")
        self.assertEqual(df.iloc[1]["variedad"], "2149")

    def test_transform_parses_fec_siem_stored_as_excel_serial_number(self) -> None:
        df_raw = self._build_raw()
        df_raw["fec_siem"] = [45807, "2026-01-21", "2026-03-11"]

        df, audit = _transform_base_sheet(df_raw)

        self.assertEqual(audit["invalid_fec_siem_rows"], 0)
        self.assertEqual(df.iloc[0]["yyww_siembra"], "2522")

    def test_palermo_and_palmas_frames_combine_with_both_fincas(self) -> None:
        df_palermo, _ = _transform_base_sheet(self._build_raw())

        df_palmas_raw = self._build_raw().assign(finca=[1.0, 1.0, None])
        df_palmas, _ = _transform_base_sheet(df_palmas_raw)

        combined = pd.concat([df_palermo, df_palmas], ignore_index=True)

        self.assertEqual(len(combined), 4)
        self.assertEqual(set(combined["finca"]), {"001", "002"})

    def _build_raw_lotes(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "Id": [12695, 12670, 99999],
                "Finca": ["PALERMO", "PALERMO", "INVERPALMAS"],
                "Bloque": [1, 1, 5],
                "Variedad": ["FARIDA", "DON PEDRO", "OTRA"],
                "Semana de siembra": [2611, 2610, 2601],
            }
        )

    def test_transform_lotes_siembra_normalizes_join_keys(self) -> None:
        df = _transform_lotes_siembra(self._build_raw_lotes())

        self.assertEqual(list(df.columns), ["finca", "bloque", "nvari", "yyww_siembra", "id_lote_siembra"])
        row = df.iloc[0]
        self.assertEqual(row["finca"], "002")
        self.assertEqual(row["bloque"], "1")
        self.assertEqual(row["nvari"], "FARIDA")
        self.assertEqual(row["yyww_siembra"], "2611")
        self.assertEqual(row["id_lote_siembra"], 12695)

    def test_transform_lotes_siembra_raises_when_column_missing(self) -> None:
        df_raw = self._build_raw_lotes().drop(columns=["Semana de siembra"])
        with self.assertRaises(ValueError):
            _transform_lotes_siembra(df_raw)

    def test_attach_id_lote_siembra_matches_on_finca_bloque_nvari_semana(self) -> None:
        df_planos, _ = _transform_base_sheet(self._build_raw())
        df_lotes = _transform_lotes_siembra(self._build_raw_lotes())

        result = _attach_id_lote_siembra(df_planos, df_lotes)

        self.assertIn("id_lote_siembra", result.columns)
        self.assertEqual(len(result), len(df_planos))
        matched = result[result["nvari"] == "FARIDA"]
        self.assertEqual(matched.iloc[0]["id_lote_siembra"], 12695)

    def test_attach_id_lote_siembra_leaves_null_when_no_match(self) -> None:
        df_planos, _ = _transform_base_sheet(self._build_raw())
        df_lotes = _transform_lotes_siembra(self._build_raw_lotes()).iloc[0:0]

        result = _attach_id_lote_siembra(df_planos, df_lotes)
        self.assertTrue(result["id_lote_siembra"].isna().all())

    def test_attach_id_lote_siembra_does_not_duplicate_rows_on_ambiguous_key(self) -> None:
        df_planos, _ = _transform_base_sheet(self._build_raw())
        df_lotes = pd.concat(
            [
                _transform_lotes_siembra(self._build_raw_lotes()),
                _transform_lotes_siembra(self._build_raw_lotes().assign(Id=[1, 2, 3])),
            ],
            ignore_index=True,
        )

        result = _attach_id_lote_siembra(df_planos, df_lotes)
        self.assertEqual(len(result), len(df_planos))

    def _build_matched(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "finca": ["002", "002", "001", "001"],
                "bloque": ["1", "1", "5", "5"],
                "variedad": ["2048", "2149", "2308", "2431"],
                "cosecha": ["SE2638", "JL2631", "AB2616", "JL2629"],
                "nvari": ["FARIDA", "MOON_LIGHT", "NOVIA", "UKELELE"],
                "yyww_siembra": ["2611", "2604", "2539", "2501"],
                # sin_cambios: Factor=1, Mueve semanas=0 -> se excluye (nada cambio).
                # solo_mueve_semanas: Factor=1, Mueve semanas!=0 -> se incluye (OR).
                # solo_factor: Factor!=1, Mueve semanas=0 -> se incluye (OR).
                # ambos_cambian: Factor!=1 y Mueve semanas!=0 -> se incluye.
                "rev": [0.0, 2.0, 0.0, -3.0],
                "amort": [1.0, 1.0, 0.5, 1.2],
                "proyecta_quipus": [1.0, 1.0, 1.0, 1.0],
                "id_lote_siembra": [111, 222, 333, 444],
            }
        )

    def _build_datos_df(self) -> pd.DataFrame:
        # variedad lleva guion bajo (como en datos.xlsx). NOVIA tiene
        # quipus-proyección vacio (dato que no viene de Khipus, se descarta).
        # UKELELE si tiene quipus-proyección pero no tiene nombre coincidente
        # en ld_variedades (ver _build_variedad_lookup), tambien se descarta.
        # finca queda como int (2/1), igual que en datos.xlsx real: prueba
        # que _build_hoja1_export la normaliza a 3 digitos (002/001).
        return pd.DataFrame(
            {
                "finca": [2, 2, 1, 1],
                "flor": ["CLA", "CM0", "CLA", "CLA"],
                "bloque": ["1", "1", "5", "5"],
                "cosecha": ["SE2638", "JL2631", "AB2616", "JL2629"],
                "variedad": ["FARIDA", "MOON_KING", "NOVIA", "UKELELE"],
                "sem_prod": [2611, 2604, 2539, 2611],
                "inverpalmas-proy-inicial": [10, 20, 30, 40],
                "inverpalmas-proyección": [11, 21, 31, 41],
                "inverpalmas-real": [12, 22, None, 42],
                "quipus-proyección": [500, 600, None, 800],
                "color": ["ROJO", "ROJO", "BLANCO", "BLANCO"],
                "fiesta": ["", "", "", ""],
                "nfiesta": ["", "", "", ""],
            }
        )

    def _build_variedad_lookup(self) -> pd.DataFrame:
        # codigo distinto del texto de "variedad" a proposito, para probar
        # que Hoja1 exporta el codigo real (necesario para subir al sistema
        # de cultivo), no el nombre. nombre lleva espacios (como en ld_variedades) mientras
        # variedad en datos.xlsx lleva guion bajo. "UKELELE" no tiene
        # nombre coincidente aqui.
        return pd.DataFrame(
            {
                "codigo": ["9048", "9149", "9308"],
                "nombre": ["FARIDA", "MOON KING", "NOVIA"],
            }
        )

    def test_build_ajustes_export_keeps_rows_where_either_factor_or_mueve_semanas_changed(self) -> None:
        result = _build_ajustes_export(self._build_matched())

        self.assertEqual(set(result["ID"]), {222, 333, 444})
        self.assertNotIn(111, set(result["ID"]))

    def test_build_ajustes_export_renames_and_adds_fixed_columns(self) -> None:
        result = _build_ajustes_export(self._build_matched())

        self.assertEqual(list(result.columns), ["finca", "ID", "Factor", "Mueve semanas", "Vigencia", "Bloqueado"])
        row = result[result["ID"] == 444].iloc[0]
        self.assertEqual(row["finca"], "001")
        self.assertEqual(row["Factor"], 1.2)
        self.assertEqual(row["Mueve semanas"], -3.0)
        self.assertEqual(row["Vigencia"], 10)
        self.assertEqual(row["Bloqueado"], 0)

    def test_build_ajustes_export_keeps_finca_for_filtering_but_not_in_excel_columns(self) -> None:
        result = _build_ajustes_export(self._build_matched())

        self.assertIn("finca", result.columns)
        excel_columns = [c for c in result.columns if c != "finca"]
        self.assertEqual(excel_columns, ["ID", "Factor", "Mueve semanas", "Vigencia", "Bloqueado"])

    def test_build_ajustes_export_nan_factor_does_not_falsely_qualify_unchanged_row(self) -> None:
        # id 111 no tiene cambios reales (Factor=1, Mueve semanas=0); si el
        # guard de notna() no funcionara, pandas evalua "NaN != 1" como True
        # y la fila se colaria por error.
        df = self._build_matched()
        df.loc[df["id_lote_siembra"] == 111, "amort"] = float("nan")

        result = _build_ajustes_export(df)

        self.assertNotIn(111, set(result["ID"]))

    def test_load_base_from_staging_returns_rows_for_finca(self) -> None:
        fake_rows = [("002", "1", "2048", "SE2638", "FARIDA", "2026-03-11", 0.0, 1.0, 0.88)]
        with patch(
            "app.services.etl_jobs.planos_proyeccion_generator.execute_query", return_value=fake_rows
        ) as mock_query:
            df = _load_base_from_staging("002")

        self.assertEqual(list(df.columns), SOURCE_COLUMNS)
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0]["nvari"], "FARIDA")
        self.assertIn("finca = '002'", mock_query.call_args[0][0])
        self.assertIn("estado = 'activo'", mock_query.call_args[0][0])

    def test_load_base_from_staging_returns_empty_df_when_table_missing(self) -> None:
        error = mysql.connector.errors.ProgrammingError("Table 'informes.stg_plano_proyeccion' doesn't exist")
        error.errno = 1146
        with patch(
            "app.services.etl_jobs.planos_proyeccion_generator.execute_query", side_effect=error
        ):
            df = _load_base_from_staging("001")

        self.assertEqual(list(df.columns), SOURCE_COLUMNS)
        self.assertTrue(df.empty)

    def _build_ajustes_for_files(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "finca": ["002", "001"],
                "ID": [111, 222],
                "Factor": [1.2, 0.8],
                "Mueve semanas": [-3.0, 2.0],
                "Vigencia": [10, 10],
                "Bloqueado": [0, 0],
            }
        )

    def test_write_ajustes_por_finca_writes_one_file_per_finca_without_sending_email(self) -> None:
        import tempfile

        df_ajustes = self._build_ajustes_for_files()
        df_missing = pd.DataFrame(columns=["finca", "bloque", "variedad", "cosecha", "nvari", "yyww_siembra"])

        with tempfile.TemporaryDirectory() as tmp_dir, patch(
            "app.services.etl_jobs.planos_proyeccion_generator.send_email"
        ) as mock_send:
            resultados = _write_ajustes_por_finca(df_ajustes, df_missing, Path(tmp_dir))

            mock_send.assert_not_called()
            for info in resultados.values():
                self.assertTrue(Path(info["output"]).exists())

        self.assertEqual(resultados["002"]["ajustes_rows"], 1)
        self.assertEqual(resultados["001"]["ajustes_rows"], 1)

    def test_build_hoja2_export_keeps_only_rows_with_real_adjustment(self) -> None:
        result = _build_hoja2_export(self._build_matched())

        self.assertEqual(list(result.columns), ["finca", "bloque", "variedad", "cosecha", "rev", "amort"])
        self.assertEqual(len(result), 3)
        self.assertEqual(set(result["variedad"]), {"2149", "2308", "2431"})

    def test_build_hoja1_export_uses_hoja1_headers_and_derives_año_semana_from_sem_prod(self) -> None:
        result = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO)

        self.assertEqual(list(result.columns), HOJA1_COLUMNS)
        # NOVIA (sin quipus-proyección) y UKELELE (sin nombre en ld_variedades) se excluyen.
        self.assertEqual(len(result), 2)

        row = result.iloc[0]
        self.assertEqual(row["tipo"], "AJUSTE")
        self.assertEqual(row["año"], 2026)
        self.assertEqual(row["semana"], "11")
        self.assertEqual(row["finca"], "002")
        self.assertEqual(row["flor"], "CLA")
        self.assertEqual(row["bloque"], 1)
        self.assertEqual(row["cosecha"], "SE2638")
        # variedad exporta el codigo de ld_variedades ("9048"), no el nombre
        # ("FARIDA") de datos.xlsx: la plantilla necesita el codigo real.
        self.assertEqual(row["variedad"], "9048")
        # tallos viene de quipus-proyección, no de inverpalmas-real/proyección.
        self.assertEqual(row["tallos"], 500)

    def test_build_hoja1_export_derives_año_4_digitos_y_semana_2_digitos_from_sem_prod(self) -> None:
        result = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO)

        row = result.iloc[1]
        self.assertEqual(row["año"], 2026)
        self.assertEqual(row["semana"], "04")

    def test_build_hoja1_export_uses_given_tipo(self) -> None:
        result = _build_hoja1_export(
            self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO, tipo="INICIAL"
        )
        self.assertEqual(set(result["tipo"]), {"INICIAL"})

    def test_generate_planos_proyeccion_hoja1_rejects_invalid_tipo(self) -> None:
        with self.assertRaises(ValueError):
            generate_planos_proyeccion_hoja1({"tipo": "OTRO"})

    def test_build_hoja1_export_keeps_only_weeks_from_desde(self) -> None:
        # Corte en la semana 2611: se descarta MOON_KING (2604).
        result = _build_hoja1_export(
            self._build_datos_df(), self._build_variedad_lookup(), desde=yyww_to_date(2611)
        )
        self.assertEqual(result["variedad"].tolist(), ["9048"])
        self.assertEqual(result["semana"].tolist(), ["11"])

    def test_build_hoja1_export_defaults_desde_to_next_week(self) -> None:
        with patch(
            "app.services.etl_jobs.planos_proyeccion_generator.get_next_week_thursday",
            return_value=yyww_to_date(2605),
        ):
            result = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup())
        self.assertEqual(result["variedad"].tolist(), ["9048"])

    def test_yyww_series_to_date_returns_thursday_and_nat_for_invalid(self) -> None:
        result = yyww_series_to_date(pd.Series([2611, "2604", 2604.0, None, "abc", 2699]))
        self.assertEqual(result.iloc[0], pd.Timestamp("2026-03-12"))
        self.assertEqual(result.iloc[1], pd.Timestamp("2026-01-22"))
        self.assertEqual(result.iloc[2], pd.Timestamp("2026-01-22"))
        self.assertTrue(result.iloc[3:].isna().all())

    def test_get_next_week_thursday_is_thursday_of_following_week(self) -> None:
        import datetime as dt

        # Miercoles 2026-09-16 (semana 2638) -> jueves de la 2639.
        with patch(
            "app.services.utils.date_service.get_bogota_now",
            return_value=dt.datetime(2026, 9, 16, 10, 0),
        ):
            self.assertEqual(get_next_week_thursday(), yyww_to_date(2639))
        # Domingo 2026-09-20 sigue siendo semana 2638.
        with patch(
            "app.services.utils.date_service.get_bogota_now",
            return_value=dt.datetime(2026, 9, 20, 23, 0),
        ):
            self.assertEqual(get_next_week_thursday(), yyww_to_date(2639))

    def test_build_hoja1_export_drops_rows_with_tallos_zero_or_negative(self) -> None:
        df_datos = self._build_datos_df()
        df_datos.loc[df_datos["variedad"] == "MOON_KING", "quipus-proyección"] = 0
        df_datos = pd.concat(
            [df_datos, df_datos[df_datos["variedad"] == "FARIDA"].assign(**{"quipus-proyección": -5})],
            ignore_index=True,
        )
        result = _build_hoja1_export(df_datos, self._build_variedad_lookup(), desde=DESDE_TODO)
        self.assertEqual(result["variedad"].tolist(), ["9048"])
        self.assertEqual(result["tallos"].tolist(), [500])

    def test_build_hoja1_export_drops_rows_without_quipus_proyeccion(self) -> None:
        result = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO)
        # NOVIA (codigo "9308" en el lookup) no tiene dato en quipus-proyección.
        self.assertNotIn("9308", result["variedad"].tolist())

    def test_build_hoja1_export_drops_rows_whose_variedad_has_no_match_in_ld_variedades(self) -> None:
        result = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO)
        # UKELELE si tiene quipus-proyección pero no aparece en el lookup.
        self.assertEqual(set(result["variedad"]), {"9048", "9149"})

    def test_build_hoja1_export_raises_when_variedad_lookup_is_empty(self) -> None:
        with self.assertRaises(RuntimeError):
            _build_hoja1_export(
                self._build_datos_df(), pd.DataFrame(columns=["codigo", "nombre"]), desde=DESDE_TODO
            )

    def test_build_hoja1_export_raises_when_required_column_missing(self) -> None:
        df_datos = self._build_datos_df().drop(columns=["quipus-proyección"])
        with self.assertRaises(ValueError):
            _build_hoja1_export(df_datos, self._build_variedad_lookup(), desde=DESDE_TODO)

    def test_load_quipus_excel_raises_when_file_missing(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            missing_path = Path(tmp_dir) / "no_existe.xlsx"
            with patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.QUIPUS_EXCEL_OUTPUT_PATH",
                str(missing_path),
            ):
                with self.assertRaises(FileNotFoundError):
                    _load_quipus_excel()

    def test_load_quipus_excel_reads_hoja_datos(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "datos.xlsx"
            self._build_datos_df().to_excel(excel_path, index=False, sheet_name="datos")

            with patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.QUIPUS_EXCEL_OUTPUT_PATH",
                str(excel_path),
            ):
                df = _load_quipus_excel()

            self.assertEqual(len(df), 4)
            self.assertIn("quipus-proyección", df.columns)

    def test_generate_planos_proyeccion_hoja1_fills_template_hoja1_and_hoja2_and_sends_it(self) -> None:
        import tempfile

        import openpyxl

        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "datos.xlsx"
            self._build_datos_df().to_excel(excel_path, index=False, sheet_name="datos")

            csv_path = Path(tmp_dir) / "planos_proyeccion.csv"
            self._build_matched().to_csv(csv_path, index=False)

            template_path = Path(tmp_dir) / "template.xlsx"
            self._write_template(template_path)

            out_path = Path(tmp_dir) / "planos_proyeccion_hoja1.xlsx"
            with patch(
                "app.core.settings.settings.PLANOS_PROYECCION_HOJA1_OUTPUT_PATH", str(out_path)
            ), patch(
                "app.core.settings.settings.QUIPUS_EXCEL_OUTPUT_PATH", str(excel_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_OUTPUT_PATH", str(csv_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_TIKA_TEMPLATE_PATH", str(template_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_CULTIVO_EMAILS", ["estadistica@inverpalmas.com"]
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_CULTIVO_CC_EMAILS", []
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator._load_variedad_codigos",
                return_value=self._build_variedad_lookup(),
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.get_next_week_thursday",
                return_value=DESDE_TODO,
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.send_email"
            ) as mock_send:
                result_path, audit = generate_planos_proyeccion_hoja1({"tipo": "inicial"}, debug=True)

            self.assertEqual(str(Path(result_path)), str(out_path))
            self.assertEqual(audit["tipo"], "INICIAL")
            self.assertEqual(audit["datos_rows"], 4)
            self.assertEqual(audit["hoja1_rows"], 2)
            self.assertEqual(audit["hoja2_rows"], 3)
            self.assertTrue(out_path.exists())

            mock_send.assert_called_once()
            _, kwargs = mock_send.call_args
            self.assertEqual(kwargs["to"], ["estadistica@inverpalmas.com"])
            self.assertIsNone(kwargs["cc"])
            self.assertEqual(kwargs["attachments"], [out_path])
            self.assertIn(out_path.name, kwargs["body"])
            self.assertIn("sistema de cultivo", kwargs["subject"])
            self.assertNotIn("Tika", kwargs["subject"] + kwargs["body"])

            wb = openpyxl.load_workbook(out_path)
            headers = [cell.value for cell in wb["Hoja1"][1]]
            self.assertEqual(headers, HOJA1_COLUMNS)
            hoja1_rows = list(wb["Hoja1"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(hoja1_rows), 2)
            self.assertEqual(hoja1_rows[0][0], "INICIAL")
            self.assertEqual(hoja1_rows[0][3], "002")
            self.assertEqual(hoja1_rows[0][5], 1)
            self.assertIsInstance(hoja1_rows[0][5], int)
            self.assertEqual(hoja1_rows[0][7], "9048")

            hoja2_rows = list(wb["Hoja2"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(hoja2_rows), 3)
            self.assertNotIn("999", [row[0] for row in hoja2_rows])
            self.assertEqual({row[2] for row in hoja2_rows}, {"2149", "2308", "2431"})
            self.assertTrue(all(isinstance(row[1], int) for row in hoja2_rows))

    def test_generate_planos_proyeccion_hoja1_raises_when_template_missing(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir, patch(
            "app.core.settings.settings.PLANOS_PROYECCION_TIKA_TEMPLATE_PATH",
            str(Path(tmp_dir) / "no_existe.xlsx"),
        ):
            with self.assertRaises(FileNotFoundError):
                generate_planos_proyeccion_hoja1()

    def test_build_tika_upload_body_mentions_upload_url_steps_and_ajustes_files(self) -> None:
        body = _build_tika_upload_body(["ajustes_proyeccion_palermo.xlsx", "ajustes_proyeccion_palmas.xlsx"])
        self.assertIn("inverpalmas-tika.meliusid.app/khipus/projection_adjustments", body)
        self.assertNotIn("plantilla", body.lower())
        self.assertIn("ajustes_proyeccion_palermo.xlsx", body)
        self.assertIn("ajustes_proyeccion_palmas.xlsx", body)

    def _write_template(self, path: Path) -> None:
        import openpyxl

        wb = openpyxl.Workbook()
        hoja1 = wb.active
        hoja1.title = "Hoja1"
        hoja1.append(["tipo", "año", "semana", "finca", "flor", "bloque", "cosecha", "variedad", "tallos"])
        hoja1.append(["x", 2026, 1, "001", "ROS", "1", "EN2601", "2048", 100])

        hoja2 = wb.create_sheet("Hoja2")
        hoja2.append(["finca", "bloque", "variedad", "cosecha", "rev", "amort"])
        hoja2.append(["999", "99", "9999", "OLD9999", -9, 0.1])
        wb.save(path)

    def test_fill_subir_proyeccion_workbook_writes_hoja1_and_hoja2_replacing_template_rows(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            template_path = Path(tmp_dir) / "template.xlsx"
            self._write_template(template_path)

            df_hoja1 = _build_hoja1_export(self._build_datos_df(), self._build_variedad_lookup(), desde=DESDE_TODO)
            df_hoja2 = _build_hoja2_export(self._build_matched())
            output_path = Path(tmp_dir) / "output.xlsx"
            _fill_subir_proyeccion_workbook(df_hoja1, df_hoja2, template_path, output_path)

            import openpyxl

            wb = openpyxl.load_workbook(output_path)
            hoja1_rows = list(wb["Hoja1"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(hoja1_rows), len(df_hoja1))
            self.assertNotIn("x", [row[0] for row in hoja1_rows])
            hoja2_rows = list(wb["Hoja2"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(hoja2_rows), len(df_hoja2))
            self.assertNotIn("999", [row[0] for row in hoja2_rows])

    def _write_ajustes_files(self, tmp_dir: str) -> dict[str, dict[str, object]]:
        import openpyxl

        ajustes_files: dict[str, dict[str, object]] = {}
        for finca_code, finca_name in (("002", "palermo"), ("001", "palmas")):
            path = Path(tmp_dir) / f"ajustes_proyeccion_{finca_name}.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["ID", "Factor", "Mueve semanas", "Vigencia", "Bloqueado"])
            wb.save(path)
            ajustes_files[finca_code] = {"output": str(path), "ajustes_rows": 1, "missing_rows": 0}
        return ajustes_files

    def test_notify_tika_ajuste_sends_only_ajustes_files(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            ajustes_files = self._write_ajustes_files(tmp_dir)

            with patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.PLANOS_PROYECCION_TIKA_EMAILS",
                ["estadistica@inverpalmas.com"],
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.PLANOS_PROYECCION_TIKA_CC_EMAILS",
                [],
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.send_email"
            ) as mock_send:
                resultado = _notify_tika_ajuste(ajustes_files)

            mock_send.assert_called_once()
            _, kwargs = mock_send.call_args
            self.assertEqual(kwargs["to"], ["estadistica@inverpalmas.com"])
            self.assertIsNone(kwargs["cc"])
            self.assertEqual(len(kwargs["attachments"]), 2)
            for attachment in kwargs["attachments"]:
                self.assertTrue(attachment.name.startswith("ajustes_proyeccion_"))
            self.assertEqual(len(resultado["ajustes_files"]), 2)

    def test_notify_tika_ajuste_skips_send_when_ajustes_file_missing(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            ajustes_files = {
                "002": {"output": str(Path(tmp_dir) / "no_existe.xlsx")},
            }

            with patch(
                "app.services.etl_jobs.planos_proyeccion_generator.send_email"
            ) as mock_send:
                resultado = _notify_tika_ajuste(ajustes_files)

            mock_send.assert_not_called()
            self.assertIsNone(resultado)

    def test_notify_tika_ajuste_attaches_consolidado_with_finca_files(self) -> None:
        import tempfile

        import openpyxl

        with tempfile.TemporaryDirectory() as tmp_dir:
            ajustes_files = self._write_ajustes_files(tmp_dir)
            consolidado = Path(tmp_dir) / "ajustes_proyeccion.xlsx"
            openpyxl.Workbook().save(consolidado)

            with patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.PLANOS_PROYECCION_TIKA_EMAILS",
                ["estadistica@inverpalmas.com"],
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.settings.PLANOS_PROYECCION_TIKA_CC_EMAILS",
                [],
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.send_email"
            ) as mock_send:
                resultado = _notify_tika_ajuste(ajustes_files, consolidado)

            _, kwargs = mock_send.call_args
            nombres = [attachment.name for attachment in kwargs["attachments"]]
            self.assertEqual(len(nombres), 3)
            self.assertIn("ajustes_proyeccion.xlsx", nombres)
            self.assertIn("ajustes_proyeccion.xlsx", kwargs["body"])
            self.assertEqual(resultado["consolidado"], str(consolidado))

    def test_hoja1_week_range_returns_first_and_last_week_as_yyww(self) -> None:
        df_hoja1 = pd.DataFrame(
            {"año": [2026, 2027, 2026], "semana": ["39", "05", "52"]}
        )

        self.assertEqual(_hoja1_week_range(df_hoja1), ("2639", "2705"))

    def test_hoja1_week_range_returns_none_when_empty(self) -> None:
        self.assertIsNone(_hoja1_week_range(pd.DataFrame(columns=["año", "semana"])))

    def test_subir_proyeccion_body_warns_to_delete_projection_by_tipo_and_flor(self) -> None:
        body = _build_subir_proyeccion_body("Subir_Proyeccion.xlsx", 10, 3, "AJUSTE", ("2639", "2705"))

        self.assertIn("Antes de subir", body)
        self.assertIn("borra la proyección", body)
        self.assertIn("CLA", body)
        self.assertIn("CM0", body)
        # El tipo aparece tanto como tipo de carga como en el aviso de borrado.
        self.assertGreaterEqual(body.count("AJUSTE"), 2)

    def test_subir_proyeccion_body_reports_week_range_and_next_step(self) -> None:
        body = _build_subir_proyeccion_body("Subir_Proyeccion.xlsx", 10, 3, "INICIAL", ("2639", "2705"))

        self.assertIn("2639", body)
        self.assertIn("2705", body)
        self.assertIn("semana siguiente a la actual", body)
        self.assertIn("proyecciones_csv", body)

    def test_subir_proyeccion_body_omits_week_range_when_unknown(self) -> None:
        body = _build_subir_proyeccion_body("Subir_Proyeccion.xlsx", 0, 0, "AJUSTE", None)

        self.assertNotIn("La proyección va desde", body)
        self.assertIn("CM0", body)

    def test_generate_planos_proyeccion_hoja1_email_carries_tipo_and_week_range(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "datos.xlsx"
            self._build_datos_df().to_excel(excel_path, index=False, sheet_name="datos")

            csv_path = Path(tmp_dir) / "planos_proyeccion.csv"
            self._build_matched().to_csv(csv_path, index=False)

            template_path = Path(tmp_dir) / "template.xlsx"
            self._write_template(template_path)

            out_path = Path(tmp_dir) / "Subir_Proyeccion.xlsx"
            with patch(
                "app.core.settings.settings.PLANOS_PROYECCION_HOJA1_OUTPUT_PATH", str(out_path)
            ), patch(
                "app.core.settings.settings.QUIPUS_EXCEL_OUTPUT_PATH", str(excel_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_OUTPUT_PATH", str(csv_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_TIKA_TEMPLATE_PATH", str(template_path)
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_CULTIVO_EMAILS", ["estadistica@inverpalmas.com"]
            ), patch(
                "app.core.settings.settings.PLANOS_PROYECCION_CULTIVO_CC_EMAILS", []
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator._load_variedad_codigos",
                return_value=self._build_variedad_lookup(),
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.get_next_week_thursday",
                return_value=DESDE_TODO,
            ), patch(
                "app.services.etl_jobs.planos_proyeccion_generator.send_email"
            ) as mock_send:
                _, audit = generate_planos_proyeccion_hoja1({"tipo": "AJUSTE"}, debug=True)

            _, kwargs = mock_send.call_args
            self.assertIn("AJUSTE", kwargs["subject"])
            self.assertIn("CLA", kwargs["body"])
            self.assertIn("CM0", kwargs["body"])
            self.assertEqual(audit["semana_inicio"], min(audit["semana_inicio"], audit["semana_fin"]))
            self.assertIn(audit["semana_inicio"], kwargs["body"])


if __name__ == "__main__":
    unittest.main()
