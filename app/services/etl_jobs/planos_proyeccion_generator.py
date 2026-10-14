"""Extrae el plano de siembra/proyeccion de Palermo y Palmas desde MySQL.

Cada finca sube su plano ("BASE PLANO PROYECCION") desde una macro de
Excel a la tabla de staging STAGING_TABLE_NAME (via el endpoint PHP
plano_proyeccion_upload.php) en vez de que este job lea el .xlsm
directamente -- asi ninguna finca depende de que un archivo local se
sincronice hasta este servidor. Cada subida nueva no borra la anterior:
la marca como estado='historico' y deja la nueva como 'activo', asi se
conserva el historico y este job siempre toma solo la ultima subida
vigente por finca. Este job lee esa tabla para cada finca
(FINCA_CODES), une los resultados, cruza con Lotes_de_siembra.csv
(finca/bloque/variedad/semana de siembra) para traer el Id del lote, y
exporta un CSV con las columnas: finca, bloque, variedad, cosecha,
nvari, yyww_siembra, rev, amort, proyecta_quipus, id_lote_siembra.

yyww_siembra es el codigo semana ISO (ej. 2604) derivado de fec_siem; la
fecha de siembra completa no se conserva. nvari reemplaza espacios por
guion bajo. variedad se lee y exporta como texto (nunca se convierte a
numero) para no perder ceros a la izquierda ni caracteres no numericos.

Las filas que no cruzan con Lotes_de_siembra.csv (sin id_lote_siembra) se
excluyen de planos_proyeccion.csv. Se genera ademas ajustes_proyeccion.xlsx
en Outputs (combinado) con las filas cruzadas cuyo Factor (amort) sea
distinto de 1 o cuyo Mueve semanas (rev) sea distinto de 0, en columnas
ID/Factor/Mueve semanas/Vigencia/Bloqueado, y un Excel por finca
(ajustes_proyeccion_<finca>.xlsx). Ya no se envia correo por finca (la
tarea esta unificada); solo un correo con los ajustes para cargar a Tika.
La plantilla Subir_datos_de_proyecciones.xlsx (Hoja1 + Hoja2 llenas) para
el sistema de cultivo la genera generate_planos_proyeccion_hoja1.

Las filas cruzadas tambien se replican en MySQL (tabla
ld_planos_proyeccion), siguiendo el mismo patron que ld_proyecciones: la
API PHP (api/endpoints/planos_proyeccion.php) puede leerla directo sin
que este servicio FastAPI exponga HTTP hacia afuera.
"""
from __future__ import annotations

import datetime
import logging
from pathlib import Path

import mysql.connector
import openpyxl
import pandas as pd

from app.core.mail import MailNotConfiguredError, send_email
from app.core.settings import settings
from app.services.utils.date_service import get_next_week_thursday, yyww_series_to_date
from app.services.utils.dbf_reader import normalize_path
from app.services.etl.helpers import dataframe_to_sql, execute_query, normalize_finca_series, validate_columns

logger = logging.getLogger(__name__)

MYSQL_TABLE_NAME = "ld_planos_proyeccion"
STAGING_TABLE_NAME = "stg_plano_proyeccion"
_ERRNO_TABLE_DOES_NOT_EXIST = 1146

# Codigo de finca -> se recorren en este orden al construir el CSV/tabla final.
FINCA_CODES = ["002", "001"]  # 002 = Palermo, 001 = Palmas
FINCA_NAMES = {"002": "Palermo", "001": "Palmas"}

EXCEL_SERIAL_DATE_ORIGIN = pd.Timestamp("1899-12-30")

AJUSTES_MYSQL_TABLE_NAME = "ld_ajustes_proyeccion"

AJUSTES_COLUMN_MAP = {
    "id_lote_siembra": "ID",
    "amort": "Factor",
    "rev": "Mueve semanas",
}
AJUSTES_OUTPUT_COLUMNS = ["ID", "Factor", "Mueve semanas", "Vigencia", "Bloqueado"]
# "finca" se conserva junto a las columnas de salida solo para poder
# filtrar por finca en MySQL (api/endpoints/ajustes_proyeccion.php); el
# Excel que se descarga NO incluye "finca", solo AJUSTES_OUTPUT_COLUMNS.
AJUSTES_EXPORT_COLUMNS = ["finca"] + AJUSTES_OUTPUT_COLUMNS
AJUSTES_SQL_COLUMN_MAP = {
    "ID": "id_lote_siembra",
    "Factor": "factor",
    "Mueve semanas": "mueve_semanas",
    "Vigencia": "vigencia",
    "Bloqueado": "bloqueado",
}

# planos_proyeccion envia a Tika los ajustes por finca (TIKA_EMAIL_SUBJECT);
# planos_proyeccion_hoja1 genera y envia, para el sistema de cultivo, la
# plantilla Subir_datos_de_proyecciones.xlsx con Hoja1 y Hoja2 llenas
# (SUBIR_PROYECCION_EMAIL_SUBJECT). Hoja2 usa el mismo filtro de "ajuste"
# (Factor!=1 o Mueve semanas!=0) que AJUSTES_*, pero conservando
# bloque/variedad/cosecha en vez de renombrar a ID/Factor/etc.
TIKA_EMAIL_SUBJECT = "para Subir a Tika - Ajuste de proyección"
SUBIR_PROYECCION_EMAIL_SUBJECT = "para Subir al sistema de cultivo - Proyección"
# Columnas de planos_proyeccion.csv que se leen como texto (ceros a la izquierda).
PLANOS_CSV_TEXT_COLUMNS = ["finca", "bloque", "variedad", "cosecha", "nvari", "yyww_siembra"]
HOJA2_COLUMNS = ["finca", "bloque", "variedad", "cosecha", "rev", "amort"]
TIKA_UPLOAD_URL = "https://inverpalmas-tika.meliusid.app/khipus/projection_adjustments"

# Encabezados de la Hoja1 de la plantilla Subir_datos_de_proyecciones.xlsx.
# Se arma a partir de la hoja "datos" de datos.xlsx (job quipus_vs_proy,
# settings.QUIPUS_EXCEL_OUTPUT_PATH): tipo es configurable (AJUSTE/INICIAL); año (4
# digitos) y semana (2 digitos) se derivan de sem_prod (YYWW); flor ya viene
# como codflor en esa hoja (CLA/CM0, ver settings.QUIPUS_FLOWER_MAP); tallos
# es la columna quipus-proyección (la unica cantidad que viene de Khipus en
# ese archivo -- las demas son proyeccion/real internas de Inverpalmas);
# variedad exporta el codigo de ld_variedades (no el nombre), porque la
# plantilla necesita el codigo real. Ver _build_hoja1_export.
# tipo es parametro del job planos_proyeccion_hoja1 (por defecto AJUSTE).
HOJA1_TIPO_DEFAULT = "AJUSTE"
HOJA1_TIPOS_VALIDOS = ("AJUSTE", "INICIAL")
HOJA1_COLUMNS = ["tipo", "año", "semana", "finca", "flor", "bloque", "cosecha", "variedad", "tallos"]

QUIPUS_EXCEL_SHEET_NAME = "datos"
QUIPUS_EXCEL_CANTIDAD_COLUMN = "quipus-proyección"

# Columnas de la hoja "datos" de datos.xlsx que _build_hoja1_export necesita.
QUIPUS_EXCEL_REQUIRED_COLUMNS = {
    "finca",
    "flor",
    "bloque",
    "cosecha",
    "variedad",
    "sem_prod",
    QUIPUS_EXCEL_CANTIDAD_COLUMN,
}

LOTES_REQUIRED_COLUMNS = {"Id", "Finca", "Bloque", "Variedad", "Semana de siembra"}

REQUIRED_COLUMNS = {
    "finca",
    "bloque",
    "variedad",
    "cosecha",
    "nvari",
    "fec_siem",
    "rev",
    "amort",
    "proyecta_quipus",
}

SOURCE_COLUMNS = [
    "finca",
    "bloque",
    "variedad",
    "cosecha",
    "nvari",
    "fec_siem",
    "rev",
    "amort",
    "proyecta_quipus",
]

OUTPUT_COLUMNS = [
    "finca",
    "bloque",
    "variedad",
    "cosecha",
    "nvari",
    "yyww_siembra",
    "rev",
    "amort",
    "proyecta_quipus",
]

JOIN_KEYS = ["finca", "bloque", "nvari", "yyww_siembra"]

FINAL_COLUMNS = OUTPUT_COLUMNS + ["id_lote_siembra"]


def _load_base_from_staging(finca_code: str) -> pd.DataFrame:
    """Lee las filas subidas por la macro de Excel para una finca.

    Solo lee estado='activo': cada subida nueva marca el lote anterior de
    esa finca como 'historico' en vez de borrarlo (ver
    plano_proyeccion_upload.php), asi que aqui siempre se toma unicamente
    la ultima subida vigente por finca.

    finca_code viene de la lista fija FINCA_CODES (no de entrada de
    usuario), asi que interpolarlo directo en el SQL es seguro.
    """
    columns = list(SOURCE_COLUMNS)
    sql = (
        f"SELECT {', '.join(columns)} FROM {STAGING_TABLE_NAME} "
        f"WHERE finca = '{finca_code}' AND estado = 'activo'"
    )
    try:
        rows = execute_query(sql) or []
    except mysql.connector.errors.ProgrammingError as exc:
        if exc.errno == _ERRNO_TABLE_DOES_NOT_EXIST:
            logger.warning(
                "planos_proyeccion: la tabla %s aun no existe; ninguna macro ha subido datos todavia.",
                STAGING_TABLE_NAME,
            )
            return pd.DataFrame(columns=columns)
        raise

    return pd.DataFrame(rows, columns=columns)


def _parse_fec_siem(value: object) -> pd.Timestamp:
    """Convierte una celda fec_siem a Timestamp.

    Algunas filas del Excel guardan la fecha como numero serial (celda con
    formato General/Numero en vez de Fecha); pandas/openpyxl las entrega
    como int/float en lugar de datetime, y pd.to_datetime las interpretaria
    mal (como nanosegundos desde 1970). Se detectan y convierten usando el
    origen de fechas de Excel (1899-12-30).
    """
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return pd.NaT
    if isinstance(value, (pd.Timestamp, datetime.datetime, datetime.date)):
        return pd.Timestamp(value)
    if isinstance(value, (int, float)):
        return EXCEL_SERIAL_DATE_ORIGIN + pd.to_timedelta(float(value), unit="D")
    return pd.to_datetime(value, errors="coerce", dayfirst=True)


def _transform_base_sheet(df: pd.DataFrame, sheet_name: str = "BASE PLANO PROYECCION") -> tuple[pd.DataFrame, dict[str, int]]:
    audit: dict[str, object] = {"raw_rows": int(len(df))}

    validate_columns(df, REQUIRED_COLUMNS, sheet_name)
    df = df[SOURCE_COLUMNS].copy()

    df = df[df["finca"].notna()].copy()
    audit["rows_after_dropna_finca"] = int(len(df))

    df["finca"] = normalize_finca_series(df["finca"])
    df["bloque"] = pd.to_numeric(df["bloque"], errors="coerce").astype("Int64").astype(str)
    df["variedad"] = df["variedad"].astype(str).str.strip()
    df["cosecha"] = df["cosecha"].astype(str).str.strip()
    df["nvari"] = df["nvari"].astype(str).str.strip().str.replace(r"\s+", "_", regex=True)

    fec_siem = pd.to_datetime(df["fec_siem"].apply(_parse_fec_siem))
    audit["invalid_fec_siem_rows"] = int(fec_siem.isna().sum())
    iso_week = fec_siem.dt.isocalendar()["week"].astype("Int64").astype(str).str.zfill(2)
    df["yyww_siembra"] = fec_siem.dt.strftime("%y") + iso_week
    df = df.drop(columns=["fec_siem"])

    df["rev"] = pd.to_numeric(df["rev"], errors="coerce")
    df["amort"] = pd.to_numeric(df["amort"], errors="coerce")
    df["proyecta_quipus"] = pd.to_numeric(df["proyecta_quipus"], errors="coerce")

    df = df[OUTPUT_COLUMNS].reset_index(drop=True)
    audit["final_rows"] = int(len(df))

    return df, audit


def _load_lotes_siembra(input_path: str) -> pd.DataFrame:
    path = Path(normalize_path(input_path))
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de lotes de siembra: {path}")

    df = pd.read_csv(path, encoding="utf-8-sig")
    df.columns = [str(col).strip() for col in df.columns]
    return df


def _transform_lotes_siembra(df: pd.DataFrame) -> pd.DataFrame:
    missing = LOTES_REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"Faltan columnas en Lotes_de_siembra.csv: {', '.join(sorted(missing))}")

    df = df[["Id", "Finca", "Bloque", "Variedad", "Semana de siembra"]].copy()
    df["finca"] = normalize_finca_series(df["Finca"])
    df["bloque"] = pd.to_numeric(df["Bloque"], errors="coerce").astype("Int64").astype(str)
    df["nvari"] = df["Variedad"].astype(str).str.strip().str.replace(r"\s+", "_", regex=True)
    df["yyww_siembra"] = pd.to_numeric(df["Semana de siembra"], errors="coerce").astype("Int64").astype(str).str.zfill(4)
    df["id_lote_siembra"] = pd.to_numeric(df["Id"], errors="coerce").astype("Int64")

    return df[JOIN_KEYS + ["id_lote_siembra"]]


def _attach_id_lote_siembra(df_planos: pd.DataFrame, df_lotes: pd.DataFrame) -> pd.DataFrame:
    duplicated_keys = df_lotes.duplicated(subset=JOIN_KEYS, keep=False)
    df_lotes_dedup = df_lotes[~duplicated_keys].drop_duplicates(subset=JOIN_KEYS)

    merged = df_planos.merge(df_lotes_dedup, on=JOIN_KEYS, how="left")
    return merged[FINAL_COLUMNS]


def _ajustes_mask(df_matched: pd.DataFrame) -> pd.Series:
    """Fila cambio si Factor (amort) != 1 o Mueve semanas (rev) != 0."""
    factor = df_matched["amort"]
    mueve_semanas = df_matched["rev"]
    return (factor.notna() & (factor != 1)) | (mueve_semanas.notna() & (mueve_semanas != 0))


def _build_ajustes_export(df_matched: pd.DataFrame) -> pd.DataFrame:
    """Filtra Factor!=1 o Mueve semanas!=0, y arma las columnas finales.

    Incluye "finca" ademas de AJUSTES_OUTPUT_COLUMNS para poder filtrar por
    finca al persistir en MySQL y al escribir el Excel por finca; quien
    escriba el Excel a descargar/enviar debe quitarla antes (ver
    AJUSTES_OUTPUT_COLUMNS).
    """
    mask = _ajustes_mask(df_matched)
    df_ajustes = df_matched[mask].rename(columns=AJUSTES_COLUMN_MAP).copy()
    df_ajustes["Vigencia"] = 10
    df_ajustes["Bloqueado"] = 0
    return df_ajustes[AJUSTES_EXPORT_COLUMNS].reset_index(drop=True)


def _build_hoja2_export(df_matched: pd.DataFrame) -> pd.DataFrame:
    """Filas de ajuste (mismo filtro que _build_ajustes_export) en forma de
    Hoja2 de Subir_datos_de_proyecciones.xlsx: finca/bloque/variedad/cosecha/
    rev/amort, ambas fincas juntas (identificadas por la columna finca)."""
    mask = _ajustes_mask(df_matched)
    out = df_matched[mask][HOJA2_COLUMNS].copy()
    # bloque se exporta como entero (igual que en Hoja1), no como texto.
    out["bloque"] = pd.to_numeric(out["bloque"], errors="coerce").astype("Int64")
    return out.reset_index(drop=True)


def _load_quipus_excel() -> pd.DataFrame:
    """Lee la hoja "datos" de datos.xlsx, generado semanalmente por el job
    quipus_vs_proy (settings.QUIPUS_EXCEL_OUTPUT_PATH)."""
    excel_path = Path(normalize_path(settings.QUIPUS_EXCEL_OUTPUT_PATH))
    if not excel_path.exists():
        raise FileNotFoundError(
            f"No existe {excel_path}; corre primero el job 'quipus_vs_proy' de esta semana."
        )
    return pd.read_excel(excel_path, sheet_name=QUIPUS_EXCEL_SHEET_NAME)


def _load_variedad_codigos() -> pd.DataFrame:
    """Mapeo codigo/nombre de ld_variedades, para resolver el codigo de variedad en Hoja1."""
    rows = execute_query("SELECT codigo, nombre FROM ld_variedades") or []
    return pd.DataFrame(rows, columns=["codigo", "nombre"])


def _build_hoja1_export(
    df_datos: pd.DataFrame,
    df_variedad: pd.DataFrame,
    desde: pd.Timestamp | None = None,
    tipo: str = HOJA1_TIPO_DEFAULT,
) -> pd.DataFrame:
    """Hoja1 de la plantilla armada a partir de la hoja "datos" de datos.xlsx.

    Usa los encabezados de Hoja1 de Subir_datos_de_proyecciones.xlsx
    (HOJA1_COLUMNS). Solo se exportan las filas con quipus-proyección
    (tallos) mayor a cero: las demas columnas de cantidad de esa hoja
    son proyeccion/real internas de Inverpalmas, no de Khipus. finca se
    exporta en 3 digitos (001/002, regla fija para exportaciones a Excel);
    año (4 digitos) y semana (2 digitos) se derivan de sem_prod (YYWW);
    flor ya viene como codflor en datos.xlsx; variedad exporta el codigo de
    ld_variedades (no el nombre) porque la plantilla necesita el codigo
    real -- se cruza variedad (nombre con guion bajo) contra
    ld_variedades.nombre (con espacios), igual que en quipus.csv; las filas
    sin cruce se excluyen.

    Solo se exportan semanas desde la siguiente a la actual: sem_prod (YYWW)
    se convierte al jueves de su semana y se compara contra `desde` (por
    defecto get_next_week_thursday()). tipo va igual en todas las filas.
    """
    if df_variedad.empty:
        raise RuntimeError("No se pudo cargar ld_variedades para resolver el codigo de variedad en Hoja1.")

    validate_columns(df_datos, QUIPUS_EXCEL_REQUIRED_COLUMNS, "datos.xlsx (hoja datos)")

    if desde is None:
        desde = get_next_week_thursday()

    fecha_semana = yyww_series_to_date(df_datos["sem_prod"])
    tallos = pd.to_numeric(df_datos[QUIPUS_EXCEL_CANTIDAD_COLUMN], errors="coerce")
    out = df_datos[(tallos > 0) & (fecha_semana >= desde)].copy()

    out["finca"] = normalize_finca_series(out["finca"])
    # bloque se exporta como entero (no texto) para que el Excel no lo marque como numero-en-texto.
    out["bloque"] = pd.to_numeric(out["bloque"], errors="coerce").astype("Int64")

    sem_prod = pd.to_numeric(out["sem_prod"], errors="coerce").astype("Int64").astype(str).str.zfill(4)
    out["tipo"] = tipo
    out["año"] = 2000 + sem_prod.str[:2].astype(int)
    out["semana"] = sem_prod.str[2:4]
    out["tallos"] = pd.to_numeric(out[QUIPUS_EXCEL_CANTIDAD_COLUMN], errors="coerce")

    out["variedad_nombre"] = out["variedad"].astype(str).str.strip().str.replace("_", " ", regex=False)
    df_variedad = df_variedad.copy()
    df_variedad["nombre"] = df_variedad["nombre"].astype(str).str.strip()
    out = out.merge(df_variedad[["codigo", "nombre"]], how="inner", left_on="variedad_nombre", right_on="nombre")
    out["variedad"] = out["codigo"]

    return out[HOJA1_COLUMNS].reset_index(drop=True)


def _write_ajustes_por_finca(df_ajustes: pd.DataFrame, df_missing: pd.DataFrame, ajustes_dir: Path) -> dict[str, object]:
    """Escribe un Excel de ajustes por finca (ajustes_proyeccion_<finca>.xlsx).

    Ya no se envia correo por finca: la tarea esta unificada y estos
    archivos solo viajan adjuntos en el correo de Tika (_notify_tika_ajuste).
    """
    resultados: dict[str, object] = {}

    for finca_code in FINCA_CODES:
        finca_name = FINCA_NAMES[finca_code]
        df_ajustes_finca = df_ajustes[df_ajustes["finca"] == finca_code][AJUSTES_OUTPUT_COLUMNS].reset_index(
            drop=True
        )
        df_missing_finca = df_missing[df_missing["finca"] == finca_code].reset_index(drop=True)

        finca_output_file = ajustes_dir / f"ajustes_proyeccion_{finca_name.lower()}.xlsx"
        with pd.ExcelWriter(finca_output_file, engine="openpyxl") as writer:
            df_ajustes_finca.to_excel(writer, index=False, sheet_name="ajustes")

        resultados[finca_code] = {
            "output": str(finca_output_file),
            "ajustes_rows": int(len(df_ajustes_finca)),
            "missing_rows": int(len(df_missing_finca)),
        }

    return resultados


def _write_sheet_rows(sheet, df: pd.DataFrame, columns: list[str]) -> None:
    """Borra las filas de datos (deja el encabezado) y escribe df[columns]."""
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in df[columns].astype(object).itertuples(index=False):
        sheet.append([None if pd.isna(value) else value for value in row])


def _fill_subir_proyeccion_workbook(
    df_hoja1: pd.DataFrame, df_hoja2: pd.DataFrame, template_path: Path, output_path: Path
) -> Path:
    """Copia la plantilla Subir_datos_de_proyecciones.xlsx con Hoja1 (datos
    de quipus-proyección) y Hoja2 (filas de ajuste) ya diligenciadas."""
    wb = openpyxl.load_workbook(template_path)
    _write_sheet_rows(wb["Hoja1"], df_hoja1, HOJA1_COLUMNS)
    _write_sheet_rows(wb["Hoja2"], df_hoja2, HOJA2_COLUMNS)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _build_tika_upload_body(ajustes_filenames: list[str], consolidado_filename: str | None = None) -> str:
    archivos = " y ".join(f"<strong>{name}</strong>" for name in ajustes_filenames)
    body = f"<p>Se adjuntan los archivos de ajustes por finca ({archivos})"
    if consolidado_filename:
        body += f" y el consolidado de ambas fincas (<strong>{consolidado_filename}</strong>)"
    return (
        f"{body}.</p>"
        f'<p>Subirlos a <a href="{TIKA_UPLOAD_URL}">{TIKA_UPLOAD_URL}</a> usando el importador alli '
        "dispuesto.</p>"
    )


def _hoja1_week_range(df_hoja1: pd.DataFrame) -> tuple[str, str] | None:
    """Primera y ultima semana (YYWW) presentes en la hoja de proyeccion.

    Se reporta en el correo para que el operador confirme, antes de cargar,
    que la proyeccion arranca en la semana siguiente a la actual.
    """
    if df_hoja1.empty:
        return None
    semanas = (
        (pd.to_numeric(df_hoja1["año"], errors="coerce") - 2000) * 100
        + pd.to_numeric(df_hoja1["semana"], errors="coerce")
    ).dropna()
    if semanas.empty:
        return None
    return f"{int(semanas.min()):04d}", f"{int(semanas.max()):04d}"


def _build_subir_proyeccion_body(
    filename: str, hoja1_rows: int, hoja2_rows: int, tipo: str, week_range: tuple[str, str] | None
) -> str:
    body = (
        f"<p>Se adjunta <strong>{filename}</strong> (plantilla Subir_datos_de_proyecciones.xlsx) con "
        f"sus dos hojas diligenciadas y lista para subir al sistema de cultivo: {hoja1_rows} fila(s) de "
        "proyección y "
        f"{hoja2_rows} fila(s) de ajuste.</p>"
        f"<p>Tipo de carga: <strong>{tipo}</strong>.</p>"
    )
    if week_range:
        primera, ultima = week_range
        body += (
            f"<p>La proyección va desde la semana <strong>{primera}</strong> hasta la "
            f"<strong>{ultima}</strong> (arranca en la semana siguiente a la actual). Confírmalo antes de "
            "cargar.</p>"
        )
    body += (
        "<p><strong>Antes de subir el archivo</strong>, en el sistema de cultivo borra la proyección "
        f"existente por tipo (<strong>{tipo}</strong>) y por flor, tanto <strong>CLA</strong> como "
        "<strong>CM0</strong>. Solo entonces carga el Excel adjunto.</p>"
        "<p>Ya cargado, genera en el sistema de cultivo la consulta de proyecciones y ejecuta el job "
        "<code>proyecciones_csv</code> (Proyecciones a CSV) para cerrar el ciclo.</p>"
    )
    return body


def _clean_emails(addresses: list[str] | None) -> list[str]:
    return [str(addr).strip() for addr in (addresses or []) if str(addr).strip()]


def _notify_tika_ajuste(
    ajustes_files: dict[str, dict[str, object]], consolidado_path: Path | None = None
) -> dict[str, object] | None:
    """Envia por correo los archivos de ajustes por finca (los que hay que
    subir primero a Tika) y, si se pasa consolidado_path, tambien el
    ajustes_proyeccion.xlsx consolidado de ambas fincas. La plantilla Subir_datos_de_proyecciones.xlsx ya
    no va aqui: la genera y envia generate_planos_proyeccion_hoja1.

    ajustes_files es el resultado de _write_ajustes_por_finca: por cada
    finca trae la ruta del ajustes_proyeccion_<finca>.xlsx ya escrito.

    Nunca lanza excepcion: un problema de envio queda en el log pero no
    debe tumbar el ETL.
    """
    ajustes_paths = [Path(info["output"]) for info in ajustes_files.values() if info.get("output")]
    attachments = ajustes_paths + ([consolidado_path] if consolidado_path else [])
    missing_paths = [str(path) for path in attachments if not path.exists()]
    if missing_paths:
        logger.warning(
            "planos_proyeccion: faltan archivos de ajustes por finca (%s); no se envia el correo Tika.",
            ", ".join(missing_paths),
        )
        return None

    destinatarios = _clean_emails(settings.PLANOS_PROYECCION_TIKA_EMAILS)
    if not destinatarios:
        logger.warning(
            "planos_proyeccion: PLANOS_PROYECCION_TIKA_EMAILS esta vacio; no se envia el correo Tika."
        )
        return None

    cc = _clean_emails(settings.PLANOS_PROYECCION_TIKA_CC_EMAILS)

    try:
        send_email(
            to=destinatarios,
            subject=TIKA_EMAIL_SUBJECT,
            body=_build_tika_upload_body(
                [path.name for path in ajustes_paths], consolidado_path.name if consolidado_path else None
            ),
            html=True,
            cc=cc or None,
            attachments=attachments,
        )
    except MailNotConfiguredError as exc:
        logger.warning("planos_proyeccion: no se envio el correo Tika: %s", exc)
        return None
    except Exception:  # noqa: BLE001
        logger.exception("planos_proyeccion: fallo inesperado enviando el correo Tika.")
        return None

    return {
        "ajustes_files": [str(path) for path in ajustes_paths],
        "consolidado": str(consolidado_path) if consolidado_path else None,
        "to": destinatarios,
        "cc": cc,
    }


def _notify_subir_proyeccion(
    output_file: Path,
    hoja1_rows: int,
    hoja2_rows: int,
    tipo: str,
    week_range: tuple[str, str] | None,
) -> dict[str, object] | None:
    """Envia la plantilla ya diligenciada para el sistema de cultivo. Nunca lanza excepcion: el archivo
    ya quedo escrito en Outputs de todas formas."""
    destinatarios = _clean_emails(settings.PLANOS_PROYECCION_CULTIVO_EMAILS)
    if not destinatarios:
        logger.warning(
            "planos_proyeccion_hoja1: PLANOS_PROYECCION_CULTIVO_EMAILS esta vacio; no se envia la plantilla."
        )
        return None

    cc = _clean_emails(settings.PLANOS_PROYECCION_CULTIVO_CC_EMAILS)

    try:
        send_email(
            to=destinatarios,
            subject=f"{SUBIR_PROYECCION_EMAIL_SUBJECT} ({tipo})",
            body=_build_subir_proyeccion_body(
                output_file.name, hoja1_rows, hoja2_rows, tipo, week_range
            ),
            html=True,
            cc=cc or None,
            attachments=[output_file],
        )
    except MailNotConfiguredError as exc:
        logger.warning("planos_proyeccion_hoja1: no se envio la plantilla: %s", exc)
        return None
    except Exception:  # noqa: BLE001
        logger.exception("planos_proyeccion_hoja1: fallo inesperado enviando la plantilla.")
        return None

    return {"to": destinatarios, "cc": cc}


def generate_planos_proyeccion(debug: bool = False) -> str | tuple[str, dict[str, int]]:
    output_path = settings.PLANOS_PROYECCION_OUTPUT_PATH
    lotes_path = settings.LOTES_SIEMBRA_INPUT_PATH

    frames: list[pd.DataFrame] = []
    audit: dict[str, object] = {"por_finca": {}}
    for finca_code in FINCA_CODES:
        df_raw = _load_base_from_staging(finca_code)
        if df_raw.empty:
            logger.warning(
                "planos_proyeccion: no hay filas en %s para finca %s (la macro no ha subido datos todavia).",
                STAGING_TABLE_NAME,
                finca_code,
            )
            audit["por_finca"][finca_code] = {"raw_rows": 0}
            continue

        df_file, file_audit = _transform_base_sheet(df_raw, STAGING_TABLE_NAME)
        frames.append(df_file)
        audit["por_finca"][finca_code] = file_audit

    if not frames:
        raise RuntimeError(
            f"No hay datos en {STAGING_TABLE_NAME} para ninguna finca; "
            "verifica que las macros de Excel ya subieron informacion."
        )

    df = pd.concat(frames, ignore_index=True)
    audit["final_rows"] = int(len(df))
    audit["finca_counts"] = df["finca"].value_counts(dropna=False).to_dict()

    df_lotes_raw = _load_lotes_siembra(lotes_path)
    df_lotes = _transform_lotes_siembra(df_lotes_raw)
    df = _attach_id_lote_siembra(df, df_lotes)

    matched_mask = df["id_lote_siembra"].notna()
    df_matched = df[matched_mask].reset_index(drop=True)
    df_missing = df[~matched_mask].reset_index(drop=True)
    audit["matched_id_lote_siembra"] = int(len(df_matched))
    audit["missing_id_lote_siembra"] = int(len(df_missing))

    output_file = Path(normalize_path(output_path))
    output_file.parent.mkdir(parents=True, exist_ok=True)
    df_matched.to_csv(output_file, index=False)
    dataframe_to_sql(df_matched, MYSQL_TABLE_NAME, if_exists="replace")

    df_ajustes = _build_ajustes_export(df_matched)
    dataframe_to_sql(
        df_ajustes.rename(columns=AJUSTES_SQL_COLUMN_MAP), AJUSTES_MYSQL_TABLE_NAME, if_exists="replace"
    )

    ajustes_output_file = Path(normalize_path(settings.PLANOS_PROYECCION_AJUSTES_OUTPUT_PATH))
    ajustes_output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ajustes_output_file, engine="openpyxl") as writer:
        df_ajustes[AJUSTES_OUTPUT_COLUMNS].to_excel(writer, index=False, sheet_name="ajustes")
    audit["ajustes_rows"] = int(len(df_ajustes))
    audit["ajustes_output"] = str(ajustes_output_file)

    resultados_ajustes = _write_ajustes_por_finca(df_ajustes, df_missing, ajustes_output_file.parent)
    audit["por_finca_ajustes"] = resultados_ajustes
    audit["tika_ajuste"] = _notify_tika_ajuste(resultados_ajustes, ajustes_output_file)

    if debug:
        return str(output_file), audit
    return str(output_file)


def _load_planos_proyeccion_csv() -> pd.DataFrame:
    """Lee planos_proyeccion.csv (filas ya cruzadas), escrito por el job
    planos_proyeccion, para armar la Hoja2 de la plantilla."""
    csv_path = Path(normalize_path(settings.PLANOS_PROYECCION_OUTPUT_PATH))
    if not csv_path.exists():
        raise FileNotFoundError(f"No existe {csv_path}; corre primero el job 'planos_proyeccion'.")
    df = pd.read_csv(csv_path, dtype={col: str for col in PLANOS_CSV_TEXT_COLUMNS})
    validate_columns(df, set(HOJA2_COLUMNS), "planos_proyeccion.csv")
    df["finca"] = normalize_finca_series(df["finca"])
    df["rev"] = pd.to_numeric(df["rev"], errors="coerce")
    df["amort"] = pd.to_numeric(df["amort"], errors="coerce")
    return df


def _resolve_hoja1_tipo(value: object) -> str:
    tipo = str(value or HOJA1_TIPO_DEFAULT).strip().upper()
    if tipo not in HOJA1_TIPOS_VALIDOS:
        raise ValueError(
            f"tipo invalido '{value}'; valores permitidos: {', '.join(HOJA1_TIPOS_VALIDOS)}."
        )
    return tipo


def generate_planos_proyeccion_hoja1(
    params: dict[str, object] | None = None, debug: bool = False
) -> str | tuple[str, dict[str, int]]:
    """Genera la plantilla Subir_datos_de_proyecciones.xlsx lista para subir
    al sistema de cultivo (en settings.PLANOS_PROYECCION_HOJA1_OUTPUT_PATH) y
    la envia por correo:

    - Hoja1: hoja "datos" de datos.xlsx de la semana (job quipus_vs_proy).
    - Hoja2: filas de ajuste de planos_proyeccion.csv (job planos_proyeccion).

    params["tipo"] (AJUSTE por defecto, o INICIAL) llena la columna tipo de Hoja1.
    """
    tipo = _resolve_hoja1_tipo((params or {}).get("tipo"))
    template_path = Path(normalize_path(settings.PLANOS_PROYECCION_TIKA_TEMPLATE_PATH))
    if not template_path.exists():
        raise FileNotFoundError(f"No existe la plantilla Subir_datos_de_proyecciones.xlsx en {template_path}")

    df_datos = _load_quipus_excel()
    df_variedad = _load_variedad_codigos()
    df_hoja1 = _build_hoja1_export(df_datos, df_variedad, tipo=tipo)
    df_hoja2 = _build_hoja2_export(_load_planos_proyeccion_csv())

    output_file = Path(normalize_path(settings.PLANOS_PROYECCION_HOJA1_OUTPUT_PATH))
    _fill_subir_proyeccion_workbook(df_hoja1, df_hoja2, template_path, output_file)

    week_range = _hoja1_week_range(df_hoja1)
    audit: dict[str, object] = {
        "tipo": tipo,
        "datos_rows": int(len(df_datos)),
        "hoja1_rows": int(len(df_hoja1)),
        "hoja2_rows": int(len(df_hoja2)),
        "hoja1_output": str(output_file),
        "semana_inicio": week_range[0] if week_range else None,
        "semana_fin": week_range[1] if week_range else None,
    }
    audit["subir_proyeccion"] = _notify_subir_proyeccion(
        output_file, len(df_hoja1), len(df_hoja2), tipo, week_range
    )

    if debug:
        return str(output_file), audit
    return str(output_file)
