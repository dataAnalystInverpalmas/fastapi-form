from __future__ import annotations

import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from app.core.settings import settings
from app.services.utils.dbf_reader import normalize_path

pd.set_option('display.float_format', lambda x: '%.2f' % x)

HOJA = "base"
COLUMNS = ['FINCA', 'AREA', 'GFLOR', 'FLOR', 'GRUPO', 'SUB_GRUPO', 'TIPO', 'UMEDIDA', 'AÑO', 'MES', 'Valor']
CONVERT = {'Valor': float}
HOJA_ORDEN = 'orden'
COLUMNS_ORDEN = ['sub_grupo', 'area', 'orden', 'opera']
# columna opcional: cantidad de decimales para mostrar cuando umedida=='UND'
# (ej. HECTAREAS=2). Si la hoja 'orden' aun no tiene esta columna, se usa 0
# decimales para todo (comportamiento anterior) hasta que se agregue.
COLUMNA_DECIMALES = 'decimales'
AÑO1 = 2025
AÑO2 = 2026

FILAS_A = ['finca', 'area', 'gflor', 'flor', 'grupo', 'sub_grupo', 'tipo', 'año', 'umedida']
FILAS_C = ['flor', 'grupo', 'sub_grupo', 'tipo', 'año', 'umedida']

MESES_ORDENADOS = [
    'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO',
    'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE'
]
TIPOS_AÑO = ['REAL' + str(AÑO1), 'REAL' + str(AÑO2), 'TEORICO' + str(AÑO2)]

# grupos que entran en los reportes de empresa (Informes Consolidado / Consolidado)
GRUPOS_FILTRAR = ['AREAS', 'PLANTAS', 'PRODUCCION', 'PERSONAL', 'DESECHOS', 'NOMINA', 'COMPRAS', 'TRANSPORTES']

# sub_grupos que solo tienen sentido a nivel de poscosecha/empresa (Informes
# PC, Informes Consolidado, Consolidado), no en las hojas por finca/area de
# cultivo (Palmas, Palermo, Propagacion).
EXCLUSIVOS_PC = [
    ('PRODUCCION', 'EXPORTADO'),
    ('INDICADORES', 'Cajas Exportables'),
]

# orden fijo de 'flor' para que los resumenes por color (ROSAS, que agrupa
# NOMINA/COMPRAS/TRANSPORTES sin distincion de color) queden despues del
# detalle (ROSAS COLORES, ROSAS ROJAS) en vez de antes por orden alfabetico.
ORDEN_FLOR = {
    'CLAVEL': 1,
    'MINIATURA': 2,
    'ROSAS COLORES': 3,
    'ROSAS ROJAS': 4,
    'ROSAS': 5,
}

# filas de titulo antes del encabezado de columnas (fila 1 en blanco, fila 2
# con el rotulo de la hoja); el encabezado de columnas queda en la fila 3.
TITULO_FILAS = 2
FILA_ENCABEZADO = TITULO_FILAS + 1

# rotulo que va en la fila 2 de cada hoja de oficina (columna B), tal como en
# el archivo objetivo.
ROTULOS_OFICINA = {
    'Informes Consolidado': 'JEFE PALERMO',
    'Informes Palmas': 'JEFE PALMAS',
    'Informes Palermo': 'JEFE PALERMO',
    'Informes Propagacion': 'JEFE PALMAS',
    'Informes PC': 'JEFE PALERMO',
}

INDICADORES = {
    'FMT2 PRODUCCION': lambda x: x['PRODUCCION TOTAL'] / x['HECTAREAS'] / 10000 if x['HECTAREAS'] != 0 else 0,
    'FMT2 EXPORTABLE': lambda x: x['PRODUCCION EXPORTABLE'] / x['HECTAREAS'] / 10000 if x['HECTAREAS'] != 0 else 0,
    'FMT2 EXPORTADO': lambda x: x['EXPORTADO'] / x['HECTAREAS'] / 10000 if x['HECTAREAS'] != 0 else 0,
    'Flor Expble/Planta': lambda x: x['PRODUCCION EXPORTABLE'] / x['PLANTAS SEMBRADAS'] if x['PLANTAS SEMBRADAS'] != 0 else 0,
    'Personas x Ha': lambda x: x['PERSONAS + HORAS EXTRAS'] / x['HECTAREAS'] if x['HECTAREAS'] != 0 else 0,
    'Tallos Persona Dia': lambda x: x['PRODUCCION EXPORTABLE'] / x['PERSONAS + HORAS EXTRAS'] / 26 if x['PERSONAS + HORAS EXTRAS'] != 0 else 0,
    'Flor Exptdo/Planta': lambda x: x['EXPORTADO'] / x['PLANTAS SEMBRADAS'] if x['PLANTAS SEMBRADAS'] != 0 else 0,
    'Flor Planta Sembradas': lambda x: x['PRODUCCION TOTAL'] / x['PLANTAS SEMBRADAS'] if x['PLANTAS SEMBRADAS'] != 0 else 0,
    'Flor Planta Producción': lambda x: x['PRODUCCION TOTAL'] / x['PLANTAS EN PRODUCCION'] if x['PLANTAS EN PRODUCCION'] != 0 else 0,
}

INDICADORES1 = {
    'Cajas Exportables': lambda x: x['PRODUCCION EXPORTABLE'] / 700 if x['flor'] == 'CLAVEL' else (
        x['PRODUCCION EXPORTABLE'] / 600 if x['flor'] == 'MINIATURA' else (
            x['PRODUCCION EXPORTABLE'] / 350 if x['flor'] in ['ROSAS COLORES', 'ROSAS ROJAS', 'ROSAS'] else (
                350 if x['PRODUCCION EXPORTABLE'] != 0 else 0))),
}

INDICADORES2 = {
    'Costo Tallo': lambda x: (x['NOMINA + EXTRAS + PREST/SOCIALES/ISS'] + x['COMPRAS $ GASTOS'] + x['COSTO TRANSPORTE PERSONAL'] + x['TRANSPORTE FLOR AEROPUERTO']) / x['PRODUCCION EXPORTABLE'] if x['PRODUCCION EXPORTABLE'] != 0 else 0,
    'Costo Caja': lambda x: (x['NOMINA + EXTRAS + PREST/SOCIALES/ISS'] + x['COMPRAS $ GASTOS'] + x['COSTO TRANSPORTE PERSONAL'] + x['TRANSPORTE FLOR AEROPUERTO']) / x['Cajas Exportables'] if x['Cajas Exportables'] != 0 else 0,
}


# ---------------------------------------------------------------------------
# Helpers de bajo nivel (portados de Informes/functions.py)
# ---------------------------------------------------------------------------

def _create_df(dir_, sheet, ncolumns, convert):
    df = pd.read_excel(
        dir_,
        sheet_name=sheet,
        na_values=['(en blanco)'],
        header=0, usecols=ncolumns,
        converters=convert,
    )
    return pd.DataFrame(df)


class _DataFrameTransformerIndicators:
    def __init__(self, dataframe, columns):
        self.df = dataframe.copy()
        self.columns = columns

    def transform_and_concatenate(self, indicadores):
        df_year = self.df.copy()

        pv_table2 = df_year.pivot_table(
            index=self.columns,
            columns=['sub_grupo'],
            values='valor',
            aggfunc=np.sum,
        ).reset_index()

        nuevo_dataframe = pv_table2.copy()

        for indicador, funcion in indicadores.items():
            nuevo_dataframe[indicador] = nuevo_dataframe.apply(funcion, axis=1)

        nuevo_dataframe['GRUPO'] = 'INDICADORES'

        columnas_seleccionadas = self.columns + list(indicadores.keys())
        df_final = nuevo_dataframe[columnas_seleccionadas].reset_index(drop=True)

        df_final_melted = df_final.melt(
            id_vars=self.columns,
            var_name='sub_grupo',
            value_name='valor',
        )

        df_final_melted.sort_values(by=self.columns, inplace=True)
        df_final_melted.reset_index(drop=True, inplace=True)

        df_final_melted = df_final_melted.dropna()
        df_final_melted['grupo'] = 'INDICADORES'
        df_final_melted['umedida'] = 'UND'
        df_final_melted = df_final_melted[df_final_melted.valor > 0]
        df_final_melted = df_final_melted[self.df.columns]

        return pd.concat([self.df, df_final_melted], axis=0)


class _ProcesadorDatos:
    def __init__(self, df_indicadores, df_orden, filas, meses_ordenados, mes_acumula, tipos_año):
        self.df_indicadores = df_indicadores
        self.df_orden = df_orden
        self.filas = filas
        self.meses_ordenados = meses_ordenados
        self.pv_matriz = None
        self.mes_acumula = mes_acumula
        self.tipos_año = tipos_año

    def generar_tabla_pivotante(self):
        pv_table_c = self.df_indicadores.pivot_table(
            index=self.filas,
            columns=['mes'],
            values=['valor'],
            aggfunc={np.sum},
        )

        pv_table_c = pv_table_c.reindex(columns=self.meses_ordenados, level=2)
        pv_table_c = pv_table_c.fillna(0).reset_index()
        pv_table_c.columns = self.filas + self.meses_ordenados

        pv_table_c['tipo_año'] = pv_table_c['tipo'] + pv_table_c['año'].astype('str')

        if 'area' not in self.filas:
            pv_table_c['area'] = 'Empresa'

        self.pv_matriz = pd.merge(
            pv_table_c,
            self.df_orden,
            on=['area', 'sub_grupo'],
            how='left',
        ).reset_index(drop=True)

        self.pv_matriz.sort_values(by='tipo_año', inplace=True)
        self.pv_matriz.sort_values(by=self.filas[:-4] + ['orden', 'tipo_año', 'umedida'], inplace=True)
        self.pv_matriz = self.pv_matriz[self.pv_matriz.tipo_año.isin(self.tipos_año)]
        self.pv_matriz = self.pv_matriz.reset_index(drop=True)

    def calcular_acumulado(self):
        columnas = self.pv_matriz.columns.to_list()
        indice_enero = columnas.index("ENERO")
        indice_mes_actual = columnas.index(self.mes_acumula)
        indice_diciembre = columnas.index("DICIEMBRE")

        self.pv_matriz['ACM_mes_actual'] = self.pv_matriz.iloc[:, indice_enero:indice_mes_actual + 1].sum(axis=1)

        for index, row in self.pv_matriz.iterrows():
            if row['opera'] == 'suma':
                self.pv_matriz.at[index, 'ACM_mes_actual'] = row.iloc[indice_enero:indice_mes_actual + 1].sum()
                self.pv_matriz.loc[index, 'ACM_diciembre'] = row.iloc[indice_enero:indice_diciembre + 1].sum()
            elif row['opera'] == 'promedio':
                self.pv_matriz.at[index, 'ACM_mes_actual'] = row.iloc[indice_enero:indice_mes_actual + 1].mean()
                self.pv_matriz.loc[index, 'ACM_diciembre'] = row.iloc[indice_enero:indice_diciembre + 1].mean()

    def obtener_resultados(self):
        return self.pv_matriz.copy()


class _ExcelFormatter:
    def __init__(self, wb=None, sheet_title='Reporte', header_row=1):
        self.header_row = header_row
        if wb is None:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = sheet_title
        else:
            self.wb = wb
            self.ws = self.wb.create_sheet(title=sheet_title)

    def set_title(self, row, column, text):
        cell = self.ws.cell(row=row, column=column, value=text)
        cell.font = Font(bold=True)
        return cell

    def add_data_to_sheet(self, dataframe, start_row=1):
        for r_idx, row in enumerate(dataframe_to_rows(dataframe, index=False), start_row):
            for c_idx, value in enumerate(row, 1):
                self.ws.cell(row=r_idx, column=c_idx, value=value)

    def apply_blue_font_by_condition(self, column_number, target_value, font_color, fill_color):
        fill_style = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

        for row in self.ws.iter_rows(min_row=self.header_row + 1, max_row=self.ws.max_row, min_col=column_number, max_col=column_number):
            for cell in row:
                if cell.value == target_value:
                    for cell_in_row in self.ws.iter_cols(min_col=1, max_col=self.ws.max_column, min_row=cell.row, max_row=cell.row):
                        for cell_to_format in cell_in_row:
                            cell_to_format.font = Font(color=font_color)
                            cell_to_format.fill = fill_style

    def apply_format_by_measurement_unit(self, column_number, column_apply, decimales=None):
        start_row = self.header_row + 1
        for row in self.ws.iter_rows(min_row=start_row, max_row=self.ws.max_row, min_col=column_number, max_col=column_number):
            for cell in row:
                if cell.value == "UND":
                    n_decimales = 0
                    if decimales is not None:
                        offset = cell.row - start_row
                        if 0 <= offset < len(decimales):
                            valor = decimales[offset]
                            if pd.notna(valor) and valor:
                                n_decimales = int(valor)
                    formato = '#,##0' if n_decimales <= 0 else '#,##0.' + '0' * n_decimales
                    for cell_in_row in self.ws.iter_cols(min_col=column_apply, max_col=self.ws.max_column, min_row=cell.row, max_row=cell.row):
                        for cell_to_format in cell_in_row:
                            cell_to_format.number_format = formato
                elif cell.value == "%":
                    for cell_in_row in self.ws.iter_cols(min_col=column_apply, max_col=self.ws.max_column, min_row=cell.row, max_row=cell.row):
                        for cell_to_format in cell_in_row:
                            cell_to_format.number_format = '0.0%'
                else:
                    for cell_in_row in self.ws.iter_cols(min_col=column_apply, max_col=self.ws.max_column, min_row=cell.row, max_row=cell.row):
                        for cell_to_format in cell_in_row:
                            cell_to_format.number_format = '"$"#,##0'

    def adjust_column_width(self):
        for column in self.ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except Exception:
                    pass
            self.ws.column_dimensions[column_letter].width = max_length + 2

    def align_text_center(self):
        alignment = Alignment(horizontal='center', vertical='center')
        for row in self.ws.iter_rows(min_row=self.header_row + 1, max_row=self.ws.max_row, min_col=1, max_col=self.ws.max_column):
            for cell in row:
                cell.alignment = alignment

    def set_freeze_panes(self, column, row):
        column_letter = get_column_letter(column)
        self.ws.freeze_panes = f"{column_letter}{row}"

    def paint_duplicate_rows(self, start_column, end_column, start_row):
        current_column = start_column
        while current_column <= end_column:
            current_row = start_row + 1
            while current_row <= self.ws.max_row:
                current_cell = self.ws.cell(row=current_row, column=current_column)
                reference_cell = self.ws.cell(row=current_row - 1, column=current_column)
                if current_cell.value == reference_cell.value:
                    current_cell.font = Font(color='FFFFFF')
                else:
                    current_cell.font = Font(color='000000')
                current_row += 1
            current_column += 1

    def apply_font_color_to_range(self, start_row, start_column, end_row, end_column, font_color):
        for row in self.ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_column, max_col=end_column):
            for cell in row:
                cell.font = Font(color=font_color)

    def replace_zeros_with_empty(self, start_row, start_column, end_row, end_column):
        for row in self.ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_column, max_col=end_column):
            for cell in row:
                if cell.value == 0:
                    cell.value = ''

    def compare_and_draw_line(self, start_row, comparison_column, start_column, end_column):
        black_border = Border(top=Side(style='thin', color='000000'))
        end_row = self.ws.max_row
        current_row = start_row + 1
        while current_row <= end_row:
            current_cell = self.ws.cell(row=current_row, column=comparison_column)
            previous_cell = self.ws.cell(row=current_row - 1, column=comparison_column)
            if current_cell.value != previous_cell.value:
                for col in range(start_column, end_column + 1):
                    self.ws.cell(row=current_row, column=col).border = black_border
            current_row += 1


# ---------------------------------------------------------------------------
# Carga y transformacion (portados de Informes/informes.py)
# ---------------------------------------------------------------------------

def _strip_text_columns(df, columns):
    for col in columns:
        if col in df.columns:
            df[col] = df[col].str.strip()
    return df


# ---------------------------------------------------------------------------
# Reconstruccion en Python del flujo de Power Query
# tablas<año> -> datos_<año> -> hoja 'base' de base_datos.xlsx
#
# Reemplaza la necesidad de abrir Excel y darle "Actualizar todo" en cascada
# a datos_2026.xlsx y luego a base_datos.xlsx antes de correr el reporte: el
# job lee directamente las tablas individuales del año en curso y los
# 'datos_YYYY.xlsx' de años cerrados, y replica exactamente las mismas
# transformaciones que hacia la consulta M ('base'/'nbase'/'dist' en
# Formulas/Section1.m de base_datos.xlsx).
# ---------------------------------------------------------------------------

HOJA_WIDE = 'Hoja1'
COLUMNS_WIDE = [
    'id', 'FLOR PC', 'AREA', 'FINCA', 'GFLOR', 'FLOR', 'AÑO', 'SUB_GRUPO', 'TIPO', 'GRUPO', 'UMEDIDA', 'META',
] + MESES_ORDENADOS
HOJA_DISTRIBUCIONES = 'distribuciones'
# 'flor' combinada tal como se digita en las tablas de origen -> se reparte
# entre las flores reales usando los porcentajes de la hoja 'distribuciones'.
FLORES_COMBINADAS = settings.INFORMES_JEFES_AREA_FLORES_COMBINADAS
FLORES_DISTRIBUCION = settings.INFORMES_JEFES_AREA_FLORES_DISTRIBUCION
AÑOS_EXCLUIDOS = settings.INFORMES_JEFES_AREA_ANIOS_EXCLUIDOS


def _leer_tabla_wide(path):
    """Lee la hoja ancha (id, FLOR PC, ..., ENERO..DICIEMBRE). Algunas tablas
    individuales del año en curso todavia no tienen columnas para los meses
    futuros: se leen las columnas que existan y se completan las que falten
    con NaN (equivalente a como Power Query expande cada archivo contra el
    superconjunto de columnas del archivo de ejemplo)."""
    df = pd.read_excel(path, sheet_name=HOJA_WIDE, na_values=['(en blanco)'])
    return df.reindex(columns=COLUMNS_WIDE)


def _combinar_tablas_actuales(tablas_dir):
    """Replica la consulta 'tablas' (Folder.Files + combinar): una fila por
    archivo de tablas<año>, ignorando archivos de bloqueo de Excel (~$...)."""
    archivos = [
        os.path.join(tablas_dir, nombre)
        for nombre in os.listdir(tablas_dir)
        if nombre.lower().endswith('.xlsx') and not nombre.startswith('~$')
    ]
    if not archivos:
        raise RuntimeError(f"No se encontraron archivos .xlsx en {tablas_dir}")
    return pd.concat([_leer_tabla_wide(a) for a in archivos], ignore_index=True)


def _construir_base_desde_fuentes(tablas_dir, anios_congelados_paths, dist_df):
    frames = [_leer_tabla_wide(p) for p in anios_congelados_paths]
    frames.append(_combinar_tablas_actuales(tablas_dir))
    origen = pd.concat(frames, ignore_index=True)

    origen = origen.drop(columns=['id', 'FLOR PC', 'META'])
    # AÑO puede llegar como float (columnas Excel con celdas vacias en algunas
    # tablas individuales); si queda en float, 'REAL' + str(año) da 'REAL2025.0'
    # en vez de 'REAL2025' y el filtro de tipos/año mas adelante no encuentra
    # ninguna fila.
    origen = origen.dropna(subset=['AÑO'])
    origen['AÑO'] = origen['AÑO'].astype(int)

    id_vars = ['FINCA', 'AREA', 'GFLOR', 'FLOR', 'GRUPO', 'SUB_GRUPO', 'TIPO', 'AÑO', 'UMEDIDA']
    melted = origen.melt(id_vars=id_vars, value_vars=MESES_ORDENADOS, var_name='MES', value_name='Valor')
    melted['Valor'] = pd.to_numeric(melted['Valor'], errors='coerce')

    melted = melted[(melted['GRUPO'] != 'INDICES') & (~melted['AÑO'].isin(AÑOS_EXCLUIDOS))]

    es_combinada = melted['FLOR'].isin(FLORES_COMBINADAS)
    principal = melted[~es_combinada].copy()
    combinadas = melted[es_combinada].copy()

    # 'nbase': reparte cada fila de flor combinada (ej. 'CLA Y ROS') entre las
    # flores reales segun los porcentajes de 'distribuciones' para esa
    # finca/area/gflor/flor-combinada.
    combinadas = combinadas.merge(dist_df, on=['FINCA', 'AREA', 'GFLOR', 'FLOR'], how='left')
    for flor in FLORES_DISTRIBUCION:
        combinadas[flor] = combinadas['Valor'] * combinadas[flor]
    combinadas = combinadas.drop(columns=['FLOR', 'Valor'])
    disgregado = combinadas.melt(
        id_vars=['FINCA', 'AREA', 'GFLOR', 'GRUPO', 'SUB_GRUPO', 'TIPO', 'AÑO', 'UMEDIDA', 'MES'],
        value_vars=FLORES_DISTRIBUCION, var_name='FLOR', value_name='Valor',
    )

    columnas_finales = COLUMNS
    resultado = pd.concat([principal[columnas_finales], disgregado[columnas_finales]], ignore_index=True)
    return resultado


def _cargar_orden(file):
    try:
        df_orden = _create_df(file, HOJA_ORDEN, COLUMNS_ORDEN + [COLUMNA_DECIMALES], CONVERT)
    except ValueError:
        # la hoja 'orden' todavia no tiene la columna de decimales
        df_orden = _create_df(file, HOJA_ORDEN, COLUMNS_ORDEN, CONVERT)
        df_orden[COLUMNA_DECIMALES] = pd.NA
    return _strip_text_columns(df_orden, ['sub_grupo', 'area', 'opera'])


def _cargar_datos(file, tablas_dir=None, anios_congelados_paths=None):
    """Si tablas_dir/anios_congelados_paths vienen dados, reconstruye 'base'
    en vivo desde las tablas individuales del año en curso + los datos_YYYY
    de años cerrados (ver _construir_base_desde_fuentes). Si no, lee la hoja
    'base' de base_datos.xlsx tal cual (comportamiento anterior, util si el
    reporte se quiere generar desde una foto ya congelada de 'base')."""
    if tablas_dir and anios_congelados_paths is not None:
        dist_df = _create_df(file, HOJA_DISTRIBUCIONES, ['FINCA', 'AREA', 'GFLOR', 'FLOR'] + FLORES_DISTRIBUCION, {})
        df_est = _construir_base_desde_fuentes(tablas_dir, anios_congelados_paths, dist_df)
    else:
        df_est = _create_df(file, HOJA, COLUMNS, CONVERT)

    df_est.columns = df_est.columns.str.lower()
    df_est = _strip_text_columns(df_est, ['finca', 'area', 'gflor', 'flor', 'grupo', 'sub_grupo', 'tipo', 'umedida'])

    df_orden = _cargar_orden(file)

    return df_est, df_orden


def _ajustar_flor_grupos(df_est):
    df_est = df_est.query(" valor>0 ")
    df_est.loc[(df_est['umedida'] == '$') & (df_est['gflor'] == 'MINIATURA') & (df_est['finca'] != 'PC'), 'flor'] = 'CLAVEL'
    df_est.loc[(df_est['umedida'] == '$') & ((df_est['flor'] == 'MINIATURA') | (df_est['flor'] == 'ROSAS ROJAS') | (df_est['flor'] == 'ROSAS COLORES')), 'flor'] = df_est['gflor']
    return df_est


def _agregar_produccion_pc(df_est):
    """PC (poscosecha) no reporta PRODUCCION/AREAS por si sola: se sintetiza
    como la suma de todas las fincas, para que su reporte muestre el total
    que recibio poscosecha."""
    df_prod_pc = (
        df_est.query("grupo=='PRODUCCION' or grupo=='AREAS'")
        .groupby(['gflor', 'flor', 'grupo', 'sub_grupo', 'tipo', 'año', 'umedida', 'mes'])
        .agg({'valor': 'sum'}).reset_index()
    )
    df_prod_pc[['finca', 'area']] = 'PC'
    return pd.concat([df_est, df_prod_pc], ignore_index=True)


def _construir_indicadores_completos():
    indicadores3 = INDICADORES.copy()
    indicadores3.update(INDICADORES1)
    indicadores3.update(INDICADORES2)
    return indicadores3


def _filtrar_grupos_empresa(df):
    """Grupos que componen los reportes de empresa: PRODUCCION/AREAS solo se
    toman de las fincas reales (evita duplicar el 'PC' sintetico), y
    CALIDADES solo existe a nivel de poscosecha (area == 'PC')."""
    filtro_area_calidades = (df['grupo'] == 'CALIDADES') & (df['area'] == 'PC')
    filtro_grupos = df['grupo'].isin(GRUPOS_FILTRAR)
    filtro_excluir_pc = (df['area'] == 'PC') & (df['grupo'].isin(['PRODUCCION', 'AREAS']))
    return df[(filtro_grupos | filtro_area_calidades) & ~filtro_excluir_pc]


def _pivotear(df_indicadores, df_orden, filas, tipos_año, mes_acumula, sort_cols, drop_cols):
    procesador = _ProcesadorDatos(df_indicadores, df_orden, filas, MESES_ORDENADOS, mes_acumula, tipos_año)
    procesador.generar_tabla_pivotante()
    procesador.calcular_acumulado()
    resultado = procesador.obtener_resultados()

    if 'flor' in sort_cols:
        resultado['_orden_flor'] = resultado['flor'].map(ORDEN_FLOR).fillna(99)
        sort_cols = ['_orden_flor' if col == 'flor' else col for col in sort_cols]

    if 'gflor' in sort_cols:
        # en 'Informes Consolidado' gflor se rehace igual a flor (ver
        # _generar_tabla_empresa), asi que ordenar por el texto crudo de
        # gflor descuadra el orden fijo de ORDEN_FLOR (p.ej. 'ROSAS' antes
        # que 'ROSAS COLORES' alfabeticamente). Se usa el mismo orden fijo
        # aqui para que gflor nunca reordene por delante de _orden_flor.
        resultado['_orden_gflor'] = resultado['gflor'].map(ORDEN_FLOR).fillna(99)
        sort_cols = ['_orden_gflor' if col == 'gflor' else col for col in sort_cols]

    resultado.sort_values(by=sort_cols, inplace=True)
    resultado.drop(
        columns=[c for c in drop_cols + ['_orden_flor', '_orden_gflor'] if c in resultado.columns],
        inplace=True,
    )
    return resultado.reset_index(drop=True)


def _excluir_exclusivos_pc(df):
    filtro = pd.Series(False, index=df.index)
    for grupo, sub_grupo in EXCLUSIVOS_PC:
        filtro |= (df['grupo'] == grupo) & (df['sub_grupo'] == sub_grupo)
    return df[~filtro]


def _generar_tabla_cultivo(df_est_i, df_orden, mes_acumula):
    return _pivotear(
        df_est_i, df_orden, FILAS_A, TIPOS_AÑO, mes_acumula,
        sort_cols=['finca', 'area', 'gflor', 'flor', 'orden', 'tipo_año'],
        drop_cols=['tipo_año', 'orden', 'opera'],
    )


def _generar_tabla_empresa(df_est_i, df_orden, indicadores3, mes_acumula):
    """Hoja 'Informes Consolidado': toda la empresa, con detalle por flor.

    Se agrupa por 'flor' (no por 'gflor'): PC registra sus costos con
    gflor='TODAS'/flor='CLAVEL', mientras que Palmas/Palermo ya traen
    gflor='CLAVEL'/'ROSAS' con flor=CLAVEL/ROSAS ROJAS; agrupar por gflor
    dejaria el aporte de PC en un balde aparte en vez de sumarlo al de
    Palmas/Palermo. 'gflor' se rehace igual a 'flor' al final.
    """
    base = _filtrar_grupos_empresa(df_est_i)
    agregado = (
        base.groupby(['flor', 'grupo', 'sub_grupo', 'tipo', 'año', 'umedida', 'mes'])['valor']
        .sum().reset_index()
    )
    agregado['gflor'] = agregado['flor']
    # 'Empresa' es el area usada en la hoja 'orden' para consultar orden/opera;
    # se renombra a 'Todo' al final para que coincida con el reporte esperado.
    agregado['finca'] = 'Empresa'
    agregado['area'] = 'Empresa'

    transformer = _DataFrameTransformerIndicators(agregado, ['finca', 'area', 'gflor', 'flor', 'tipo', 'año', 'mes'])
    df_empresa_i = transformer.transform_and_concatenate(indicadores3)

    resultado = _pivotear(
        df_empresa_i, df_orden, FILAS_A, TIPOS_AÑO, mes_acumula,
        sort_cols=['finca', 'area', 'gflor', 'flor', 'orden', 'tipo_año'],
        drop_cols=['tipo_año', 'orden', 'opera'],
    )
    resultado['finca'] = 'Todo'
    resultado['area'] = 'Todo'
    return resultado


def _generar_tabla_consolidado(df_est_i, df_orden, indicadores3, mes_acumula):
    """Hoja 'Consolidado': toda la empresa agrupada solo por flor."""
    base = _filtrar_grupos_empresa(df_est_i)
    df_cons = base.groupby(['flor', 'grupo', 'sub_grupo', 'tipo', 'año', 'umedida', 'mes'])['valor'].sum().reset_index()

    transformer = _DataFrameTransformerIndicators(df_cons, ['flor', 'tipo', 'año', 'mes'])
    df_cons_i1 = transformer.transform_and_concatenate(INDICADORES1)

    transformer2 = _DataFrameTransformerIndicators(df_cons_i1, ['flor', 'tipo', 'año', 'mes'])
    df_cons_i = transformer2.transform_and_concatenate(indicadores3)

    return _pivotear(
        df_cons_i, df_orden, FILAS_C, TIPOS_AÑO, mes_acumula,
        sort_cols=['flor', 'orden', 'tipo_año'],
        drop_cols=['tipo_año', 'area', 'orden', 'opera'],
    )


def _extraer_decimales(df):
    """Saca la columna auxiliar 'decimales' (si existe) del DataFrame antes
    de escribirlo a la hoja, y la devuelve como lista alineada a las filas."""
    df = df.copy()
    if COLUMNA_DECIMALES in df.columns:
        decimales = df.pop(COLUMNA_DECIMALES).tolist()
    else:
        decimales = None
    return df, decimales


def _formatear_hoja_oficina(wb, titulo, df):
    df, decimales = _extraer_decimales(df)

    excel_formatter = _ExcelFormatter(wb=wb, sheet_title=titulo, header_row=FILA_ENCABEZADO)
    excel_formatter.set_title(2, 2, ROTULOS_OFICINA[titulo])
    excel_formatter.add_data_to_sheet(df, start_row=FILA_ENCABEZADO)

    excel_formatter.apply_blue_font_by_condition(8, AÑO1, "796F6D", 'FEF8F6')
    excel_formatter.apply_blue_font_by_condition(8, AÑO2, "131212", 'FFFFFF')
    excel_formatter.apply_blue_font_by_condition(7, "TEORICO", "0000FF", 'FFFFFF')

    excel_formatter.apply_format_by_measurement_unit(9, 10, decimales=decimales)
    excel_formatter.adjust_column_width()
    excel_formatter.align_text_center()

    excel_formatter.paint_duplicate_rows(1, 5, FILA_ENCABEZADO)
    excel_formatter.apply_font_color_to_range(FILA_ENCABEZADO, 1, FILA_ENCABEZADO, 5, "131212")

    ultima_columna = excel_formatter.ws.max_column
    excel_formatter.compare_and_draw_line(FILA_ENCABEZADO, 5, 5, ultima_columna)
    excel_formatter.compare_and_draw_line(FILA_ENCABEZADO, 6, 6, ultima_columna)

    excel_formatter.replace_zeros_with_empty(start_row=FILA_ENCABEZADO, start_column=10, end_row=excel_formatter.ws.max_row, end_column=excel_formatter.ws.max_column)

    excel_formatter.set_freeze_panes(10, 2 + TITULO_FILAS)
    return excel_formatter


def _formatear_hoja_consolidado(wb, titulo, df):
    df, decimales = _extraer_decimales(df)

    excel_formatter = _ExcelFormatter(wb=wb, sheet_title=titulo, header_row=FILA_ENCABEZADO)
    excel_formatter.set_title(2, 1, 'INFORME CONSOLIDADO POR FLOR')
    excel_formatter.add_data_to_sheet(df, start_row=FILA_ENCABEZADO)

    excel_formatter.apply_blue_font_by_condition(5, AÑO1, "796F6D", 'FEF8F6')
    excel_formatter.apply_blue_font_by_condition(5, AÑO2, "131212", 'FFFFFF')
    excel_formatter.apply_blue_font_by_condition(4, "TEORICO", "0000FF", 'FFFFFF')

    excel_formatter.apply_format_by_measurement_unit(6, 7, decimales=decimales)
    excel_formatter.adjust_column_width()
    excel_formatter.align_text_center()

    excel_formatter.set_freeze_panes(6, 2 + TITULO_FILAS)
    excel_formatter.paint_duplicate_rows(1, 4, FILA_ENCABEZADO)
    excel_formatter.apply_font_color_to_range(FILA_ENCABEZADO, 1, FILA_ENCABEZADO, 5, "131212")

    ultima_columna = excel_formatter.ws.max_column
    excel_formatter.compare_and_draw_line(FILA_ENCABEZADO, 3, 3, ultima_columna)
    excel_formatter.compare_and_draw_line(FILA_ENCABEZADO, 2, 2, ultima_columna)

    excel_formatter.replace_zeros_with_empty(start_row=FILA_ENCABEZADO, start_column=10, end_row=excel_formatter.ws.max_row, end_column=excel_formatter.ws.max_column)
    return excel_formatter


def _generate_informes_jefes_area(file: str, out_dir: str, mes_acumula: str, tablas_dir: str | None = None, anios_congelados_paths: list[str] | None = None) -> str:
    os.makedirs(out_dir, exist_ok=True)

    df_est, df_orden = _cargar_datos(file, tablas_dir=tablas_dir, anios_congelados_paths=anios_congelados_paths)
    df_est = _ajustar_flor_grupos(df_est)
    df_est = _agregar_produccion_pc(df_est)

    indicadores3 = _construir_indicadores_completos()

    transformer = _DataFrameTransformerIndicators(df_est, ['finca', 'area', 'gflor', 'flor', 'tipo', 'año', 'mes'])
    df_est_i = transformer.transform_and_concatenate(indicadores3)
    df_est_i.loc[(df_est_i['gflor'] == 'CLAVEL') & (df_est_i['flor'] == 'MINIATURA'), 'gflor'] = 'MINIATURA'

    pv_table_cultivo = _generar_tabla_cultivo(df_est_i, df_orden, mes_acumula)

    tabla_empresa = _generar_tabla_empresa(df_est_i, df_orden, indicadores3, mes_acumula)
    tabla_palmas = pv_table_cultivo[(pv_table_cultivo['finca'] == 'PALMAS') & (pv_table_cultivo['area'] == 'CULTIVO')]
    tabla_palmas = _excluir_exclusivos_pc(tabla_palmas).reset_index(drop=True)
    tabla_palermo = pv_table_cultivo[(pv_table_cultivo['finca'] == 'PALERMO') & (pv_table_cultivo['area'] == 'CULTIVO')]
    tabla_palermo = _excluir_exclusivos_pc(tabla_palermo).reset_index(drop=True)
    tabla_propagacion = pv_table_cultivo[(pv_table_cultivo['finca'] == 'PALMAS') & (pv_table_cultivo['area'] == 'PROPAGACION')]
    tabla_propagacion = _excluir_exclusivos_pc(tabla_propagacion).reset_index(drop=True)
    tabla_pc = pv_table_cultivo[(pv_table_cultivo['finca'] == 'PC') & (pv_table_cultivo['area'] == 'PC')].reset_index(drop=True)
    tabla_consolidado = _generar_tabla_consolidado(df_est_i, df_orden, indicadores3, mes_acumula)

    wb = None
    for titulo, df in [
        ('Informes Consolidado', tabla_empresa),
        ('Informes Palmas', tabla_palmas),
        ('Informes Palermo', tabla_palermo),
        ('Informes Propagacion', tabla_propagacion),
        ('Informes PC', tabla_pc),
    ]:
        formatter = _formatear_hoja_oficina(wb, titulo, df)
        wb = formatter.wb

    formatter = _formatear_hoja_consolidado(wb, 'Consolidado', tabla_consolidado)
    wb = formatter.wb

    output_path = os.path.join(out_dir, f'Informe jefes de area consolidado {AÑO2}.xlsx')
    wb.save(output_path)
    wb.close()
    return output_path


def generate_informes_jefes_area(params: dict | None = None) -> str:
    """Genera el informe consolidado 'jefes de area' (6 hojas: Informes
    Consolidado, Informes Palmas, Informes Palermo, Informes Propagacion,
    Informes PC y Consolidado).

    Por defecto reconstruye 'base' en vivo desde las tablas individuales del
    año en curso (tablas2026/*.xlsx) + los datos_YYYY.xlsx de años cerrados,
    replicando en Python el mismo flujo que antes hacia Power Query
    (tablas<año> -> datos_<año> -> hoja 'base' de base_datos.xlsx). Asi el
    reporte no depende de que alguien abra Excel y refresque las consultas
    antes de correr el job. 'orden' y 'distribuciones' se siguen leyendo de
    base_datos.xlsx (son de mantenimiento manual, no cambian mes a mes).

    params opcionales:
      - file: ruta al base_datos.xlsx (orden/distribuciones; y 'base' si use_live_source=False)
      - out_dir: carpeta de salida
      - tablas_dir: carpeta de tablas individuales del año en curso
      - use_live_source: false para leer la hoja 'base' de base_datos.xlsx tal
        cual, en vez de reconstruirla desde las fuentes (comportamiento previo)
      - mes_acumula: mes (ej. 'AGOSTO') hasta el cual se acumulan ENERO..mes
        en la columna ACM_mes_actual; se actualiza cada mes (default: settings)
    """
    params = params or {}
    file = normalize_path(params.get('file') or settings.INFORMES_JEFES_AREA_INPUT_PATH)
    out_dir = normalize_path(params.get('out_dir') or settings.INFORMES_JEFES_AREA_OUTPUT_DIR)

    mes_acumula = str(params.get('mes_acumula') or settings.INFORMES_JEFES_AREA_MES_ACUMULA).strip().upper()
    if mes_acumula not in MESES_ORDENADOS:
        raise ValueError(f"mes_acumula invalido: '{mes_acumula}'. Debe ser uno de {MESES_ORDENADOS}.")

    use_live_source = params.get('use_live_source', True)
    tablas_dir = None
    anios_congelados_paths = None
    if use_live_source:
        tablas_dir = normalize_path(params.get('tablas_dir') or settings.INFORMES_JEFES_AREA_TABLAS_DIR)
        anios_congelados_paths = [normalize_path(p) for p in settings.INFORMES_JEFES_AREA_ANIOS_CONGELADOS]

    output_path = _generate_informes_jefes_area(
        file=file, out_dir=out_dir, mes_acumula=mes_acumula,
        tablas_dir=tablas_dir, anios_congelados_paths=anios_congelados_paths,
    )
    return f"xlsx:{output_path}"
